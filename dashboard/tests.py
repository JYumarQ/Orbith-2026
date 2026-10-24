"""Pruebas del Dashboard RRHH.

Cubre `DashboardView.get_plazas()`, en particular el caso de sobrecupo
(`cant_cubierta > cant_aprobada`): antes de la corrección, `totales` se
calculaba como `cubiertas + vacantes`, que con sobrecupo queda igual a
`cubiertas` y hace que `porc_cubiertas` dé un 100% falso que esconde el
sobrecupo real. Ahora `totales` es siempre la plantilla aprobada, así que el
porcentaje de ocupación puede superar el 100% cuando corresponde.
"""

from datetime import date, timedelta

from django.contrib.contenttypes.models import ContentType
from django.test import RequestFactory, TestCase
from django.urls import reverse

from bolsa.models import Aspirante
from contratos.models import CAlta, CBaja
from dashboard.filtros import resolver_filtros, unidades_visibles
from dashboard.views import (CORTES_EDAD_DEFECTO, DashboardView,
                              _bucket_edad, _categorias_edad,
                              _cortes_edad_desde_request, _rango_semana)
from nomencladores.models import (NCargo, NGrupoEscala, NMunicipio,
                                   NProvincia, NTipoContrato,
                                   NTipoUnidadOrganizativa)
from notificaciones.models import Notificacion
from strorganizativa.models import CargoPlantilla, Departamento, UnidadOrganizativa
from usuarios.models import CustomUser


class GetPlazasTests(TestCase):

    @classmethod
    def setUpTestData(cls):
        cls.tipo_unidad = NTipoUnidadOrganizativa.objects.create(
            descripcion='UEB', es_principal=False, es_subunidad=False)
        cls.unidad = UnidadOrganizativa.objects.create(
            descripcion='UEB Central', tipo=cls.tipo_unidad, grupo_nomina=1,
            orden_informe=1)
        cls.departamento = Departamento.objects.create(
            descripcion='Contabilidad', unidad_organizativa=cls.unidad, orden_informe=1)
        cls.grupo = NGrupoEscala.objects.create(nivel='IV')
        cls.ncargo = NCargo.objects.create(
            descripcion='Especialista B', cat_ocupacional='TEC',
            grupo_escala=cls.grupo, salario_basico=4000)

    def test_sin_sobrecupo_totales_coincide_con_aprobadas(self):
        """Caso normal: cubiertas <= aprobadas. `totales` debe seguir dando el
        mismo resultado que antes de la corrección (aprobadas == cubiertas+vacantes)."""
        CargoPlantilla.objects.create(
            ncargo=self.ncargo, departamento=self.departamento,
            cant_aprobada=10, cant_cubierta=6, orden_informe=1)

        plazas = DashboardView().get_plazas()

        self.assertEqual(plazas['totales'], 10)
        self.assertEqual(plazas['cubiertas'], 6)
        self.assertEqual(plazas['vacantes'], 4)
        self.assertEqual(plazas['porc_cubiertas'], 60.0)
        self.assertEqual(plazas['porc_vacantes'], 40.0)

    def test_sobrecupo_totales_es_la_plantilla_aprobada_no_cubiertas(self):
        """Con sobrecupo (cubiertas > aprobadas), `totales` debe seguir siendo
        la plantilla aprobada, `vacantes` debe quedar en 0 (nunca negativa) y
        `porc_cubiertas` debe reflejar el sobrecupo real (>100%), no un 100%
        artificial que lo esconda."""
        CargoPlantilla.objects.create(
            ncargo=self.ncargo, departamento=self.departamento,
            cant_aprobada=5, cant_cubierta=8, orden_informe=1)

        plazas = DashboardView().get_plazas()

        self.assertEqual(plazas['totales'], 5)
        self.assertEqual(plazas['cubiertas'], 8)
        self.assertEqual(plazas['vacantes'], 0)
        self.assertEqual(plazas['porc_cubiertas'], 160.0)
        self.assertEqual(plazas['porc_vacantes'], 0.0)

    def test_sin_cargos_activos_no_divide_por_cero(self):
        plazas = DashboardView().get_plazas()

        self.assertEqual(plazas['totales'], 0)
        self.assertEqual(plazas['cubiertas'], 0)
        self.assertEqual(plazas['vacantes'], 0)
        self.assertEqual(plazas['porc_cubiertas'], 0)
        self.assertEqual(plazas['porc_vacantes'], 0)


class PiramideEdadHelpersTests(TestCase):
    """`_cortes_edad_desde_request` es todo-o-nada: si la entrada no cumple
    las tres condiciones (enteros, rango razonable, estrictamente creciente),
    se descarta COMPLETA y se usan los cortes por defecto — nunca se corrige
    parcialmente quitando solo el valor inválido."""

    factory = RequestFactory()

    def _cortes(self, cortes_qs=None):
        url = '/dashboard/' if cortes_qs is None else f'/dashboard/?cortes={cortes_qs}'
        request = self.factory.get(url)
        return _cortes_edad_desde_request(request)

    def test_sin_parametro_usa_defecto(self):
        self.assertEqual(self._cortes(None), CORTES_EDAD_DEFECTO)

    def test_cortes_validos_se_usan_tal_cual(self):
        self.assertEqual(self._cortes('20,30,45,60'), [20, 30, 45, 60])

    def test_letras_caen_a_defecto_completo(self):
        self.assertEqual(self._cortes('20,treinta,45'), CORTES_EDAD_DEFECTO)

    def test_no_creciente_cae_a_defecto_completo(self):
        self.assertEqual(self._cortes('30,20,45'), CORTES_EDAD_DEFECTO)

    def test_repetidos_caen_a_defecto_completo(self):
        self.assertEqual(self._cortes('20,20,45'), CORTES_EDAD_DEFECTO)

    def test_fuera_de_rango_cae_a_defecto_completo(self):
        self.assertEqual(self._cortes('-5,20,45'), CORTES_EDAD_DEFECTO)
        self.assertEqual(self._cortes('20,45,500'), CORTES_EDAD_DEFECTO)

    def test_categorias_se_generan_desde_los_cortes(self):
        self.assertEqual(
            _categorias_edad([25, 35, 45, 55, 65]),
            ['≤25', '26-35', '36-45', '46-55', '56-65', '66+']
        )
        self.assertEqual(
            _categorias_edad([30, 60]),
            ['≤30', '31-60', '61+']
        )

    def test_bucket_edad_cae_en_el_rango_correcto(self):
        cortes = [25, 35, 45]
        self.assertEqual(_bucket_edad(20, cortes), 0)
        self.assertEqual(_bucket_edad(25, cortes), 0)
        self.assertEqual(_bucket_edad(26, cortes), 1)
        self.assertEqual(_bucket_edad(45, cortes), 2)
        self.assertEqual(_bucket_edad(46, cortes), 3)


def _crear_provincia_municipio():
    provincia = NProvincia.objects.create(nombre='La Habana', siglas='LHA')
    municipio = NMunicipio.objects.create(nombre='Playa', provincia=provincia, siglas='PLA')
    return provincia, municipio


def _ci_para_edad(edad, seq):
    """Carnet de identidad válido (11 dígitos) cuyo dueño cumple hoy `edad`
    años exactos, siguiendo la misma codificación que `Aspirante._calcular_fecha_nacimiento`."""
    hoy = date.today()
    nacimiento = hoy.replace(year=hoy.year - edad)
    century_digit = 0 if nacimiento.year < 2000 else 6
    return f"{nacimiento.year % 100:02d}{nacimiento.month:02d}{nacimiento.day:02d}{century_digit}{seq:04d}"


class HistoricoAltasBajasTests(TestCase):
    """El histórico de Altas/Bajas del dashboard (gráfico `altasBajasAreaChart`)
    se calcula agrupando `todos_altas`/`todas_bajas` en un único recorrido
    (dashboard/views.py, sección '--- 4. GRÁFICO HISTÓRICO ---'), en vez del
    O(6·N) anterior que recorría las mismas tablas una vez por mes. Aquí se
    verifica que el resultado agregado es el mismo, incluyendo el caso ya
    documentado en el código: alguien dado de alta y de baja cuenta una vez en
    altas gracias a `CBaja.fecha_alta` (CAlta se borra al causar baja)."""

    @classmethod
    def setUpTestData(cls):
        cls.tipo_unidad = NTipoUnidadOrganizativa.objects.create(
            descripcion='UEB', es_principal=False, es_subunidad=False)
        cls.unidad = UnidadOrganizativa.objects.create(
            descripcion='UEB Central', tipo=cls.tipo_unidad, grupo_nomina=1, orden_informe=1)
        cls.departamento = Departamento.objects.create(
            descripcion='Contabilidad', unidad_organizativa=cls.unidad, orden_informe=1)
        cls.grupo = NGrupoEscala.objects.create(nivel='IV')
        cls.ncargo = NCargo.objects.create(
            descripcion='Especialista B', cat_ocupacional='TEC',
            grupo_escala=cls.grupo, salario_basico=4000)
        cls.cargo = CargoPlantilla.objects.create(
            ncargo=cls.ncargo, departamento=cls.departamento, cant_aprobada=10, orden_informe=1)
        cls.provincia, cls.municipio = _crear_provincia_municipio()
        cls.tipo_contrato = NTipoContrato.objects.create(
            descripcion='Contrato Determinado', ocupa_plaza=True, requiere_motivo=False)

        hoy = date.today()

        def _meses_atras(n):
            """Misma aritmética año/mes que usa el bucle histórico de
            `DashboardView` (mes_target = hoy.month - i, con acarreo de año) —
            el día se fija a 15 porque el gráfico solo agrupa por (año, mes)."""
            mes = hoy.month - n
            anno = hoy.year
            while mes <= 0:
                mes += 12
                anno -= 1
            return date(anno, mes, 15)

        cls.hace_2_meses = _meses_atras(2)

        def _crear_aspirante(seq, estado):
            return Aspirante.objects.create(nivel_educ='SA', 
                nombre=f'Trab{seq}', papellido='Prueba', sapellido='Test',
                doc_identidad=_ci_para_edad(30, seq), estado=estado, sexo='F', raza='BL',
                provincia=cls.provincia, municipio=cls.municipio)

        # 1) Alta activa este mes: cuenta como alta en el mes actual.
        cls.asp_activo = _crear_aspirante(1, 'ACTIVO')
        cls.alta_activa = CAlta.objects.create(
            aspirante=cls.asp_activo, no_expediente='00001', tipo=cls.tipo_contrato,
            cargo=cls.cargo, fecha_alta=hoy)

        # 2) Alguien dado de alta hace 2 meses y de baja este mes: debe contar
        # una vez como alta (hace 2 meses, vía CBaja.fecha_alta) y una vez
        # como baja (este mes), sin doble conteo.
        cls.asp_baja = _crear_aspirante(2, 'BAJA')
        cls.baja = CBaja.objects.create(
            aspirante=cls.asp_baja, no_expediente='00002', tipo=cls.tipo_contrato,
            cargo=cls.cargo, fecha_alta=cls.hace_2_meses, fecha_baja=hoy)

        cls.admin = CustomUser.objects.create_user(
            username='admin_dashboard', password='x', es_admin=True)

    def test_historico_cuenta_altas_y_bajas_por_mes_sin_duplicar(self):
        self.client.force_login(self.admin)
        respuesta = self.client.get(reverse('dashboard'))
        graficos = respuesta.context['graficos']

        # Mes actual: 1 alta (asp_activo) + 1 baja (asp_baja).
        self.assertEqual(graficos['altas'][-1], 1)
        self.assertEqual(graficos['bajas'][-1], 1)
        self.assertEqual(graficos['altas_total'], 1)
        self.assertEqual(graficos['bajas_total'], 1)

        # Hace 2 meses: 1 alta (asp_baja, vía CBaja.fecha_alta), 0 bajas.
        # La ventana histórica cubre 6 meses (índice 5 = mes actual, índice 3 =
        # hace 2 meses), mismo orden que genera el bucle de `views.py`.
        idx_hace_2_meses = 3
        self.assertEqual(graficos['altas'][idx_hace_2_meses], 1)
        self.assertEqual(graficos['bajas'][idx_hace_2_meses], 0)


class FiltrosUOTests(TestCase):
    """`dashboard/filtros.py::resolver_filtros` es el único punto de alcance
    para el Dashboard: un Moderador nunca debe poder ver datos fuera de sus
    UO asignadas, ni siquiera manipulando la querystring a mano."""

    factory = RequestFactory()

    @classmethod
    def setUpTestData(cls):
        cls.tipo_unidad = NTipoUnidadOrganizativa.objects.create(
            descripcion='UEB', es_principal=False, es_subunidad=False)
        cls.unidad_propia = UnidadOrganizativa.objects.create(
            descripcion='UEB Propia', tipo=cls.tipo_unidad, grupo_nomina=1, orden_informe=1)
        cls.unidad_ajena = UnidadOrganizativa.objects.create(
            descripcion='UEB Ajena', tipo=cls.tipo_unidad, grupo_nomina=2, orden_informe=2)
        cls.depto_propio = Departamento.objects.create(
            descripcion='Contabilidad', unidad_organizativa=cls.unidad_propia, orden_informe=1)
        cls.depto_ajeno = Departamento.objects.create(
            descripcion='Almacén', unidad_organizativa=cls.unidad_ajena, orden_informe=1)

        cls.admin = CustomUser.objects.create_user(
            username='admin_filtros', password='x', es_admin=True)
        cls.observador = CustomUser.objects.create_user(
            username='obs_filtros', password='x', es_observador=True)
        cls.moderador = CustomUser.objects.create_user(
            username='mod_filtros', password='x', es_moderador=True)
        cls.moderador.unidades.add(cls.unidad_propia)

    def _request(self, usuario, **params):
        query = '&'.join(f'{k}={v}' for k, v in params.items())
        request = self.factory.get(f'/dashboard/?{query}' if query else '/dashboard/')
        request.user = usuario
        return request

    def test_moderador_solo_ve_su_propia_uo_por_defecto(self):
        uo_ids, _ = resolver_filtros(self._request(self.moderador))
        self.assertEqual(uo_ids, {self.unidad_propia.pk})

    def test_moderador_no_puede_filtrar_uo_ajena_por_querystring(self):
        uo_ids, _ = resolver_filtros(
            self._request(self.moderador, uo=self.unidad_ajena.pk))
        # El id ajeno se descarta; el resultado es el mismo que sin selección:
        # el alcance propio del moderador, nunca la UO ajena ni el universo vacío.
        self.assertEqual(uo_ids, {self.unidad_propia.pk})
        self.assertNotIn(self.unidad_ajena.pk, uo_ids)

    def test_administrador_puede_filtrar_cualquier_uo(self):
        uo_ids, _ = resolver_filtros(
            self._request(self.admin, uo=self.unidad_ajena.pk))
        self.assertEqual(uo_ids, {self.unidad_ajena.pk})

    def test_observador_puede_filtrar_cualquier_uo(self):
        uo_ids, _ = resolver_filtros(
            self._request(self.observador, uo=self.unidad_propia.pk))
        self.assertEqual(uo_ids, {self.unidad_propia.pk})

    def test_unidades_visibles_admin_ve_todas(self):
        visibles = set(unidades_visibles(self.admin).values_list('pk', flat=True))
        self.assertEqual(visibles, {self.unidad_propia.pk, self.unidad_ajena.pk})

    def test_unidades_visibles_moderador_solo_las_asignadas(self):
        visibles = set(unidades_visibles(self.moderador).values_list('pk', flat=True))
        self.assertEqual(visibles, {self.unidad_propia.pk})

    def test_uo_count_depto_count_sin_seleccion_es_alcance_completo(self):
        # Admin sin seleccionar nada: ve las 2 UO y sus 2 departamentos.
        _, contexto = resolver_filtros(self._request(self.admin))
        self.assertEqual(contexto['uo_count'], 2)
        self.assertEqual(contexto['depto_count'], 2)

    def test_uo_count_depto_count_con_una_uo_seleccionada(self):
        _, contexto = resolver_filtros(
            self._request(self.admin, uo=self.unidad_propia.pk))
        self.assertEqual(contexto['uo_count'], 1)
        self.assertEqual(contexto['depto_count'], 1)

    def test_uo_count_depto_count_con_varias_uo_seleccionadas(self):
        query = f'uo={self.unidad_propia.pk}&uo={self.unidad_ajena.pk}'
        request = self.factory.get(f'/dashboard/?{query}')
        request.user = self.admin
        _, contexto = resolver_filtros(request)
        self.assertEqual(contexto['uo_count'], 2)
        self.assertEqual(contexto['depto_count'], 2)


class RecordatoriosContratoTests(TestCase):
    """La tarjeta Recordatorios del Dashboard reutiliza
    `notificaciones.recordatorios.detectar_contratos_por_vencer` — mismo
    horizonte de 60 días y misma exclusión de contratos "Indeterminado" que
    usa la pestaña Recordatorios de Notificaciones — para no tener dos
    criterios distintos de "cuándo un contrato está por vencer"."""

    @classmethod
    def setUpTestData(cls):
        cls.tipo_unidad = NTipoUnidadOrganizativa.objects.create(
            descripcion='UEB', es_principal=False, es_subunidad=False)
        cls.unidad = UnidadOrganizativa.objects.create(
            descripcion='UEB Central', tipo=cls.tipo_unidad, grupo_nomina=1, orden_informe=1)
        cls.departamento = Departamento.objects.create(
            descripcion='Contabilidad', unidad_organizativa=cls.unidad, orden_informe=1)
        cls.grupo = NGrupoEscala.objects.create(nivel='IV')
        cls.ncargo = NCargo.objects.create(
            descripcion='Especialista B', cat_ocupacional='TEC',
            grupo_escala=cls.grupo, salario_basico=4000)
        cls.cargo = CargoPlantilla.objects.create(
            ncargo=cls.ncargo, departamento=cls.departamento, cant_aprobada=10, orden_informe=1)
        cls.provincia, cls.municipio = _crear_provincia_municipio()
        cls.tipo_determinado = NTipoContrato.objects.create(
            descripcion='Contrato Determinado', ocupa_plaza=True, requiere_motivo=False)
        cls.tipo_indeterminado = NTipoContrato.objects.create(
            descripcion='Contrato Indeterminado', ocupa_plaza=True, requiere_motivo=False)

        hoy = date.today()

        def _crear_aspirante(seq):
            return Aspirante.objects.create(nivel_educ='SA', 
                nombre=f'Trab{seq}', papellido='Prueba', sapellido='Test',
                doc_identidad=_ci_para_edad(30, seq), estado='ACTIVO', sexo='F', raza='BL',
                provincia=cls.provincia, municipio=cls.municipio)

        # Vence en 30 días: dentro del horizonte de 60 días -> debe aparecer.
        cls.asp_por_vencer = _crear_aspirante(10)
        CAlta.objects.create(
            aspirante=cls.asp_por_vencer, no_expediente='10001', tipo=cls.tipo_determinado,
            cargo=cls.cargo, fecha_alta=hoy - timedelta(days=10), duracion=40)

        # Vence en 90 días: fuera del horizonte -> no debe aparecer.
        cls.asp_lejano = _crear_aspirante(11)
        CAlta.objects.create(
            aspirante=cls.asp_lejano, no_expediente='10002', tipo=cls.tipo_determinado,
            cargo=cls.cargo, fecha_alta=hoy, duracion=90)

        # Indeterminado, vence "pronto" en teoría -> excluido siempre.
        cls.asp_indeterminado = _crear_aspirante(12)
        CAlta.objects.create(
            aspirante=cls.asp_indeterminado, no_expediente='10003', tipo=cls.tipo_indeterminado,
            cargo=cls.cargo, fecha_alta=hoy - timedelta(days=10), duracion=40)

        cls.admin = CustomUser.objects.create_user(
            username='admin_recordatorios', password='x', es_admin=True)

    def _recordatorio(self, *, unidad, titulo='Contrato por vencer', dias=30,
                      clave=''):
        """Un recordatorio ya materializado, que es como existen hoy.

        Los recordatorios dejaron de calcularse al vuelo desde las fechas del
        contrato: los materializa `notificaciones/recordatorios.py` como filas de
        `Notificacion`, y tanto el Dashboard como la pestaña de Notificaciones
        leen de ahí. Por eso el fixture crea la notificación en vez de esperar
        que la vista la deduzca de los `CAlta`.
        """
        from datetime import timedelta

        from django.contrib.contenttypes.models import ContentType
        from django.utils import timezone

        from notificaciones.models import Notificacion

        return Notificacion.objects.create(
            titulo=titulo, mensaje='...',
            tipo=Notificacion.Tipo.RECORDATORIO,
            origen=Notificacion.Origen.SISTEMA,
            codigo_regla='REC-CONTRATO',
            estado=Notificacion.Estado.ACTIVA,
            fecha_objetivo=timezone.localdate() + timedelta(days=dias),
            content_type=ContentType.objects.get_for_model(unidad),
            object_id=str(unidad.pk),
            unidad=unidad,
            # `Notificacion` es unica por (objeto, destinatario, regla,
            # clave_extra): sin variar la clave no se pueden crear dos
            # recordatorios de la misma regla para el mismo usuario.
            clave_extra=clave,
            destinatario=self.admin)

    def test_la_tarjeta_expone_el_conteo_de_recordatorios_visibles(self):
        """El Dashboard ya no decide QUÉ es un recordatorio, solo los cuenta.

        El criterio (horizonte, exclusión de indeterminados) vive ahora en
        `notificaciones/recordatorios.py` y lo cubren sus propias pruebas; aquí
        solo se protege que el Dashboard siga exponiendo `recordatorios_total` y
        su lista, que es lo que la plantilla pinta.
        """
        self._recordatorio(unidad=self.unidad)

        self.client.force_login(self.admin)
        respuesta = self.client.get(reverse('dashboard'))

        self.assertEqual(respuesta.context['recordatorios_total'], 1)
        self.assertEqual(len(respuesta.context['recordatorios']), 1)

    def test_la_tarjeta_solo_muestra_diez_aunque_haya_mas(self):
        """`recordatorios_total` cuenta todos; la lista se corta en diez."""
        for indice in range(12):
            self._recordatorio(unidad=self.unidad, titulo=f'Recordatorio {indice}',
                               clave=str(indice))

        self.client.force_login(self.admin)
        respuesta = self.client.get(reverse('dashboard'))

        self.assertEqual(respuesta.context['recordatorios_total'], 12)
        self.assertEqual(len(respuesta.context['recordatorios']), 10)

    def test_los_recordatorios_se_acotan_por_la_unidad_del_dashboard(self):
        """Mismo alcance por UO que la tarjeta de Advertencias."""
        otra = UnidadOrganizativa.objects.create(
            descripcion='UEB Norte', tipo=self.tipo_unidad, grupo_nomina=2,
            orden_informe=2)
        self._recordatorio(unidad=self.unidad, titulo='De mi unidad', clave='a')
        self._recordatorio(unidad=otra, titulo='De la otra unidad', clave='b')

        self.client.force_login(self.admin)
        respuesta = self.client.get(reverse('dashboard'), {'uo': self.unidad.pk})

        self.assertEqual(respuesta.context['recordatorios_total'], 1)
        self.assertEqual(respuesta.context['recordatorios'][0].titulo,
                         'De mi unidad')


class AdvertenciasTests(TestCase):
    """La tarjeta Advertencias reutiliza
    `notificaciones.filtros.filtrar_advertencias` (alcance por usuario/rol,
    filtro de severidad) y además se acota por la UO del Dashboard."""

    @classmethod
    def setUpTestData(cls):
        cls.tipo_unidad = NTipoUnidadOrganizativa.objects.create(
            descripcion='UEB', es_principal=False, es_subunidad=False)
        cls.unidad = UnidadOrganizativa.objects.create(
            descripcion='UEB Central', tipo=cls.tipo_unidad, grupo_nomina=1, orden_informe=1)
        cls.otra_unidad = UnidadOrganizativa.objects.create(
            descripcion='UEB Norte', tipo=cls.tipo_unidad, grupo_nomina=2, orden_informe=2)

        cls.admin = CustomUser.objects.create_user(
            username='admin_advertencias', password='x', es_admin=True)

        ct = ContentType.objects.get_for_model(Aspirante)
        cls.advertencia_critica = Notificacion.objects.create(
            titulo='Advertencia crítica', mensaje='...', tipo=Notificacion.Tipo.ALERTA,
            estado=Notificacion.Estado.ACTIVA, severidad=10,
            content_type=ct, object_id='1', unidad=cls.unidad, destinatario=cls.admin)
        cls.advertencia_otra_uo = Notificacion.objects.create(
            titulo='Advertencia de otra UO', mensaje='...', tipo=Notificacion.Tipo.ALERTA,
            estado=Notificacion.Estado.ACTIVA, severidad=20,
            content_type=ct, object_id='2', unidad=cls.otra_unidad, destinatario=cls.admin)

    def test_advertencias_aparecen_sin_filtro(self):
        self.client.force_login(self.admin)
        respuesta = self.client.get(reverse('dashboard'))
        titulos = [a.titulo for a in respuesta.context['advertencias']]
        self.assertIn('Advertencia crítica', titulos)
        self.assertIn('Advertencia de otra UO', titulos)

    def test_advertencias_se_acotan_por_uo_seleccionada(self):
        self.client.force_login(self.admin)
        respuesta = self.client.get(reverse('dashboard'), {'uo': self.unidad.pk})
        titulos = [a.titulo for a in respuesta.context['advertencias']]
        self.assertIn('Advertencia crítica', titulos)
        self.assertNotIn('Advertencia de otra UO', titulos)

    def test_advertencias_respeta_filtro_de_severidad(self):
        self.client.force_login(self.admin)
        respuesta = self.client.get(reverse('dashboard'), {'severidad': 10})
        titulos = [a.titulo for a in respuesta.context['advertencias']]
        self.assertIn('Advertencia crítica', titulos)
        self.assertNotIn('Advertencia de otra UO', titulos)


class CumpleanosSemanaTests(TestCase):
    """Las pestañas 'Semana actual' y 'Próxima semana' de Cumpleaños usan
    `_rango_semana` (mismo criterio lunes-domingo que Informe Inteligente),
    sin tocar la lógica de 'Cercanos' (que sigue basada en `doc_identidad`,
    sin cambios)."""

    @classmethod
    def setUpTestData(cls):
        cls.tipo_unidad = NTipoUnidadOrganizativa.objects.create(
            descripcion='UEB', es_principal=False, es_subunidad=False)
        cls.unidad = UnidadOrganizativa.objects.create(
            descripcion='UEB Central', tipo=cls.tipo_unidad, grupo_nomina=1, orden_informe=1)
        cls.departamento = Departamento.objects.create(
            descripcion='Contabilidad', unidad_organizativa=cls.unidad, orden_informe=1)
        cls.grupo = NGrupoEscala.objects.create(nivel='IV')
        cls.ncargo = NCargo.objects.create(
            descripcion='Especialista B', cat_ocupacional='TEC',
            grupo_escala=cls.grupo, salario_basico=4000)
        cls.cargo = CargoPlantilla.objects.create(
            ncargo=cls.ncargo, departamento=cls.departamento, cant_aprobada=10, orden_informe=1)
        cls.provincia, cls.municipio = _crear_provincia_municipio()
        cls.tipo_contrato = NTipoContrato.objects.create(
            descripcion='Contrato Determinado', ocupa_plaza=True, requiere_motivo=False)

        lunes_actual, _domingo_actual = _rango_semana(0)
        lunes_prox, _domingo_prox = _rango_semana(1)

        def _fecha_nacimiento_para(fecha_mes_dia, edad=30):
            """Fecha de nacimiento realista (edad ~30) que cae en el mismo
            (mes, día) que `fecha_mes_dia` — el año de esta última es
            irrelevante, solo importan mes/día para el cálculo de cumpleaños."""
            anno = date.today().year - edad
            try:
                return date(anno, fecha_mes_dia.month, fecha_mes_dia.day)
            except ValueError:
                return date(anno, fecha_mes_dia.month, fecha_mes_dia.day - 1)

        def _crear_con_cumple(seq, fecha_nacimiento):
            century_digit = 0 if fecha_nacimiento.year < 2000 else 6
            ci = (f"{fecha_nacimiento.year % 100:02d}{fecha_nacimiento.month:02d}"
                  f"{fecha_nacimiento.day:02d}{century_digit}{seq:04d}")
            asp = Aspirante.objects.create(nivel_educ='SA', 
                nombre=f'Cumple{seq}', papellido='Prueba', sapellido='Test',
                doc_identidad=ci, estado='ACTIVO', sexo='F', raza='BL',
                provincia=cls.provincia, municipio=cls.municipio)
            CAlta.objects.create(
                aspirante=asp, no_expediente=f'2000{seq}', tipo=cls.tipo_contrato,
                cargo=cls.cargo, fecha_alta=date.today() - timedelta(days=365))
            return asp

        cls.asp_esta_semana = _crear_con_cumple(1, _fecha_nacimiento_para(lunes_actual))
        cls.asp_prox_semana = _crear_con_cumple(
            2, _fecha_nacimiento_para(lunes_prox + timedelta(days=3)))
        cls.asp_fuera_de_rango = _crear_con_cumple(
            3, _fecha_nacimiento_para(lunes_actual - timedelta(days=15)))

        cls.admin = CustomUser.objects.create_user(
            username='admin_cumpleanos', password='x', es_admin=True)

    def test_cumpleanos_semana_actual(self):
        self.client.force_login(self.admin)
        respuesta = self.client.get(reverse('dashboard'))
        nombres = [c['nombre'] for c in respuesta.context['cumpleanos_semana_actual']]
        self.assertTrue(any('Cumple1' in n for n in nombres))
        self.assertFalse(any('Cumple2' in n for n in nombres))
        self.assertFalse(any('Cumple3' in n for n in nombres))

    def test_cumpleanos_proxima_semana(self):
        self.client.force_login(self.admin)
        respuesta = self.client.get(reverse('dashboard'))
        nombres = [c['nombre'] for c in respuesta.context['cumpleanos_semana_prox']]
        self.assertTrue(any('Cumple2' in n for n in nombres))
        self.assertFalse(any('Cumple1' in n for n in nombres))
        self.assertFalse(any('Cumple3' in n for n in nombres))


class DashboardInteractivaTests(TestCase):
    """Dashboard interactiva (decisiones de RRHH del 12-10-2026): los enlaces de las
    tarjetas conservan las unidades elegidas, la ventana de altas y bajas de un mes,
    el desglose del salario, los cumpleaños entre dos fechas y el filtro de
    prioridad de Recordatorios."""

    @classmethod
    def setUpTestData(cls):
        tipo_unidad = NTipoUnidadOrganizativa.objects.create(
            descripcion='UEB', es_principal=False, es_subunidad=False)
        cls.unidad = UnidadOrganizativa.objects.create(
            descripcion='UEB Central', tipo=tipo_unidad, grupo_nomina=1, orden_informe=1)
        cls.otra = UnidadOrganizativa.objects.create(
            descripcion='UEB Norte', tipo=tipo_unidad, grupo_nomina=1, orden_informe=2)
        ncargo = NCargo.objects.create(
            descripcion='Especialista B', cat_ocupacional='TEC',
            grupo_escala=NGrupoEscala.objects.create(nivel='IV'), salario_basico=4000)
        cls.cargo = CargoPlantilla.objects.create(
            ncargo=ncargo, cant_aprobada=10, orden_informe=1,
            departamento=Departamento.objects.create(
                descripcion='Contabilidad', unidad_organizativa=cls.unidad, orden_informe=1))
        cls.cargo_otra = CargoPlantilla.objects.create(
            ncargo=ncargo, cant_aprobada=10, orden_informe=1,
            departamento=Departamento.objects.create(
                descripcion='Ventas', unidad_organizativa=cls.otra, orden_informe=1))
        cls.provincia, cls.municipio = _crear_provincia_municipio()
        cls.tipo = NTipoContrato.objects.create(
            descripcion='Contrato Determinado', ocupa_plaza=True, requiere_motivo=False)

        def persona(seq, nacimiento, estado='ACTIVO', sexo='F'):
            ci = f'{nacimiento:%y%m%d}0{seq:04d}'
            return Aspirante.objects.create(
                nivel_educ='SA', nombre=f'Persona{seq}', papellido='Prueba', sapellido='Test',
                doc_identidad=ci, estado=estado, sexo=sexo, raza='BL',
                provincia=cls.provincia, municipio=cls.municipio)

        cls.dic = persona(1, date(1990, 12, 31))
        cls.ene = persona(2, date(1991, 1, 10), sexo='M')
        cls.feb = persona(3, date(1990, 2, 20))
        for seq, asp, cargo, alta in ((1, cls.dic, cls.cargo, date(2026, 6, 3)),
                                      (2, cls.ene, cls.cargo, date(2026, 6, 15)),
                                      (3, cls.feb, cls.cargo_otra, date(2026, 6, 20))):
            CAlta.objects.create(aspirante=asp, no_expediente=f'0000{seq}', tipo=cls.tipo,
                                 cargo=cargo, fecha_alta=alta)
        cls.de_baja = persona(4, date(1985, 5, 5), estado='BAJA', sexo='M')
        CBaja.objects.create(aspirante=cls.de_baja, no_expediente='00004', tipo=cls.tipo,
                             cargo=cls.cargo, fecha_alta=date(2025, 1, 1),
                             fecha_baja=date(2026, 6, 10))
        cls.admin = CustomUser.objects.create_user(
            username='admin_interactiva', password='x', es_admin=True)

    def setUp(self):
        self.client.force_login(self.admin)

    def test_los_enlaces_conservan_la_unidad_elegida(self):
        respuesta = self.client.get(reverse('dashboard'), {'uo': self.unidad.pk})
        self.assertEqual(respuesta.context['qs_unidades'], f'&unidad={self.unidad.pk}')
        self.assertEqual(str(respuesta.context['unidad_unica']), str(self.unidad.pk))
        # Desde el 24-10-2026 el género abre primero la ventana (con la unidad elegida
        # en `uo`); la ventana es la que lleva a Contratos con `unidad=`.
        self.assertContains(respuesta, "abrirVentanaDe('activos', 'sexo=F')")
        sin_filtro = self.client.get(reverse('dashboard'))
        self.assertEqual(sin_filtro.context['qs_unidades'], '')
        self.assertIsNone(sin_filtro.context['unidad_unica'])

    def test_ventana_de_altas_y_bajas_del_mes(self):
        url = reverse('dashboard_altas_bajas_mes')
        respuesta = self.client.get(url, {'anno': 2026, 'mes': 6})
        filas = respuesta.context['filas']
        self.assertEqual([(f['fecha'].day, f['tipo']) for f in filas],
                         [(3, 'alta'), (10, 'baja'), (15, 'alta'), (20, 'alta')])
        self.assertEqual((respuesta.context['total_altas'], respuesta.context['total_bajas']), (3, 1))
        self.assertContains(respuesta, 'Junio 2026')
        # Con una unidad elegida, solo lo de esa unidad.
        acotada = self.client.get(url, {'anno': 2026, 'mes': 6, 'uo': self.otra.pk})
        self.assertEqual(len(acotada.context['filas']), 1)
        self.assertEqual(self.client.get(url, {'anno': 'x', 'mes': 6}).status_code, 400)

    def test_desglose_del_salario(self):
        respuesta = self.client.get(reverse('dashboard_salario'))
        self.assertEqual(respuesta.context['personas'], 3)
        _titulo, _columna, por_unidad = respuesta.context['tablas'][0]
        self.assertEqual({f['nombre']: f['personas'] for f in por_unidad},
                         {'UEB Central': 2, 'UEB Norte': 1})

    def test_cumpleanos_entre_fechas_tambien_cruzando_fin_de_anno(self):
        def nombres(desde, hasta=''):
            r = self.client.get(reverse('dashboard_cumpleanos'), {'desde': desde, 'hasta': hasta})
            return [c['nombre'] for c in r.context['cumpleanos']]
        self.assertEqual(nombres('2026-01-20', '2026-02-20'), ['Persona3 Prueba'])
        self.assertEqual(nombres('2026-12-20', '2027-01-31'), ['Persona1 Prueba', 'Persona2 Prueba'])
        self.assertEqual(nombres('2030-12-31'), ['Persona1 Prueba'])  # un solo día; el año no cuenta
        self.assertTrue(self.client.get(reverse('dashboard_cumpleanos')).context['sin_rango'])

    def test_recordatorios_por_prioridad(self):
        respuesta = self.client.get(reverse('dashboard_recordatorios'), {'prioridad': 'verde'})
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta.context['recordatorios_total'], 0)


class PlazasDashboardTests(TestCase):
    """Ventana «Plazas» de Vacantes Libres / Ocupación de Plazas (12-10-2026)."""

    @classmethod
    def setUpTestData(cls):
        tipo_unidad = NTipoUnidadOrganizativa.objects.create(
            descripcion='UEB', es_principal=False, es_subunidad=False)
        unidad = UnidadOrganizativa.objects.create(
            descripcion='UEB Central', tipo=tipo_unidad, grupo_nomina=1, orden_informe=1)
        dpto = Departamento.objects.create(descripcion='Ventas', unidad_organizativa=unidad, orden_informe=1)
        grupo = NGrupoEscala.objects.create(nivel='IV')

        def cargo(nombre, aprobadas, cubiertas, en_plantilla=True):
            c = CargoPlantilla.objects.create(
                ncargo=NCargo.objects.create(descripcion=nombre, cat_ocupacional='TEC',
                                             grupo_escala=grupo, salario_basico=4000),
                departamento=dpto, cant_aprobada=aprobadas, orden_informe=CargoPlantilla.objects.count() + 1,
                en_plantilla=en_plantilla)
            CargoPlantilla.objects.filter(pk=c.pk).update(cant_cubierta=cubiertas)
            return c
        cls.con_vacantes = cargo('Técnico', 5, 3)
        cls.lleno = cargo('Operario', 2, 2)
        cls.inactivo = cargo('Antiguo', 0, 0)
        cargo('Fuera', 3, 0, en_plantilla=False)
        cls.admin = CustomUser.objects.create_user(username='admin_plazas', password='x', es_admin=True)
        cls.observador = CustomUser.objects.create_user(
            username='obs_plazas', password='x', es_observador=True)

    def nombres(self, vista, **extra):
        self.client.force_login(self.admin)
        r = self.client.get(reverse('dashboard_plazas'), {'vista': vista, **extra})
        return [c.ncargo.descripcion for c in r.context['cargos']]

    def test_vistas_y_buscador(self):
        self.assertEqual(self.nombres('vacantes'), ['Técnico'])
        self.assertEqual(self.nombres('llenas'), ['Operario'])
        self.assertEqual(self.nombres('inactivos'), ['Antiguo'])
        self.assertEqual(sorted(self.nombres('todas')), ['Antiguo', 'Operario', 'Técnico'])
        self.assertEqual(self.nombres('todas', q='oper'), ['Operario'])

    def test_contratar_aqui_solo_con_plaza_y_permiso(self):
        url = lambda c: reverse('dashboard_plazas_ocupantes', args=[c.pk])
        self.client.force_login(self.admin)
        self.assertContains(self.client.get(url(self.con_vacantes)), 'Contratar aquí')
        self.assertNotContains(self.client.get(url(self.lleno)), 'Contratar aquí')
        self.client.force_login(self.observador)
        self.assertNotContains(self.client.get(url(self.con_vacantes)), 'Contratar aquí')


class AnnosDeServicioFiltroTests(TestCase):
    """Tarjeta Años de Servicio: filtro «De __ a __ años» y ventana «Ver más»."""

    @classmethod
    def setUpTestData(cls):
        tipo_unidad = NTipoUnidadOrganizativa.objects.create(
            descripcion='UEB', es_principal=False, es_subunidad=False)
        unidad = UnidadOrganizativa.objects.create(
            descripcion='UEB Central', tipo=tipo_unidad, grupo_nomina=1, orden_informe=1)
        cargo = CargoPlantilla.objects.create(
            ncargo=NCargo.objects.create(descripcion='Técnico', cat_ocupacional='TEC',
                                         grupo_escala=NGrupoEscala.objects.create(nivel='IV'),
                                         salario_basico=4000),
            departamento=Departamento.objects.create(
                descripcion='Ventas', unidad_organizativa=unidad, orden_informe=1),
            cant_aprobada=10, orden_informe=1)
        provincia, municipio = _crear_provincia_municipio()
        tipo = NTipoContrato.objects.create(descripcion='Determinado', ocupa_plaza=True, requiere_motivo=False)
        hoy = date.today()
        for seq, annos in enumerate((2, 4, 8, 12), start=1):
            asp = Aspirante.objects.create(
                nivel_educ='SA', nombre=f'Antig{seq}', papellido='P', sapellido='T',
                doc_identidad=f'8001011270{seq}', estado='ACTIVO', sexo='F', raza='BL',
                provincia=provincia, municipio=municipio)
            CAlta.objects.create(aspirante=asp, no_expediente=f'0095{seq}', tipo=tipo, cargo=cargo,
                                 fecha_alta=hoy, fecha_alta_entidad=date(hoy.year - annos, 1, 1))
        cls.admin = CustomUser.objects.create_user(username='admin_antig', password='x', es_admin=True)

    def annos(self, url='dashboard_antiguedad_html', **params):
        self.client.force_login(self.admin)
        r = self.client.get(reverse(url), params)
        filas = r.context['page_obj'] if 'page_obj' in r.context else r.context['filas']
        return sorted(c.antiguedad for c in filas)

    def test_rangos(self):
        self.assertEqual(self.annos(annos_desde=3, annos_hasta=5), [4])
        self.assertEqual(self.annos(annos_hasta=4), [2, 4])
        self.assertEqual(self.annos(annos_desde=10), [12])
        self.assertEqual(self.annos(annos_desde=8, annos_hasta=8), [8])
        self.assertEqual(self.annos(annos_desde='x'), [2, 4, 8, 12])

    def test_ver_mas_usa_el_mismo_filtro(self):
        self.assertEqual(self.annos('dashboard_antiguedad_detalle', annos_desde=3, annos_hasta=9), [4, 8])


class PiramideDetalleTests(DashboardInteractivaTests):
    """Cada barra de la Pirámide de Edades coincide con su ventana."""

    def test_cada_barra_coincide_con_su_ventana(self):
        datos = self.client.get(reverse('dashboard_piramide_json')).json()
        for sexo, serie in (('M', datos['edades_m']), ('F', datos['edades_f'])):
            for tramo, cantidad in enumerate(serie):
                r = self.client.get(reverse('dashboard_piramide_detalle'),
                                    {'tramo': tramo, 'sexo': sexo, 'cortes': datos['cortes']})
                self.assertEqual(len(r.context['filas']), cantidad, (sexo, tramo))
        self.assertEqual(self.client.get(reverse('dashboard_piramide_detalle'),
                                         {'tramo': 99, 'sexo': 'M'}).status_code, 400)


class VentanasDashboardTests(DashboardInteractivaTests):
    """Ventanas con «Ver más» y Excel (24-10-2026): cuentan lo mismo que su tarjeta,
    «Ver más» lleva a la lista con los mismos filtros y el Excel saca lo que se ve."""

    def ventana(self, clave, **params):
        respuesta = self.client.get(reverse('dashboard_ventana', args=[clave]), params)
        self.assertEqual(respuesta.status_code, 200)
        return respuesta

    def nombres(self, clave, **params):
        return sorted(f['nombre'] for f in self.ventana(clave, **params).context['tablas'][0]['tabla'].filas)

    def test_activos_con_los_filtros_de_contratos_y_su_ver_mas(self):
        self.assertEqual(len(self.nombres('activos')), 3)
        self.assertEqual(self.nombres('activos', sexo='M'), ['Persona2 Prueba Test'])
        respuesta = self.ventana('activos', sexo='F', uo=self.unidad.pk)
        self.assertEqual([f['nombre'] for f in respuesta.context['tablas'][0]['tabla'].filas],
                         ['Persona1 Prueba Test'])
        self.assertIn('Mujeres', respuesta.context['ventana'].titulo)
        texto, url = respuesta.context['ver_mas'][0]
        self.assertEqual(url, f"{reverse('list_contrato')}?sexo=F&unidad={self.unidad.pk}")

    def test_buscar_y_paginar(self):
        self.assertEqual(self.nombres('activos', q='persona3'), ['Persona3 Prueba Test'])
        self.assertEqual(self.nombres('activos', q='00002'), ['Persona2 Prueba Test'])
        solo = self.client.get(reverse('dashboard_ventana', args=['activos']), {'solo_tabla': 1})
        self.assertTemplateUsed(solo, 'pages/dashboard/partials/ventana_tabla.html')
        self.assertTemplateNotUsed(solo, 'pages/dashboard/partials/ventana.html')

    def test_bajas_por_causa_y_su_lista(self):
        self.assertEqual(self.nombres('causa', causa='ninguna'), ['Persona4 Prueba Test'])
        self.assertEqual(self.client.get(reverse('dashboard_ventana', args=['causa']),
                                         {'causa': 'x'}).status_code, 404)
        lista = self.client.get(reverse('list_baja'), {'causa': 'ninguna'})
        self.assertEqual([a.nombre for a in lista.context['aspirantes']], ['Persona4'])

    def test_unidades_con_sus_plazas(self):
        filas = {f['nombre']: f for f in self.ventana('unidades').context['tablas'][0]['tabla'].filas}
        self.assertEqual((filas['UEB Central']['aprobadas'], filas['UEB Central']['trabajadores']), (10, 2))
        self.assertIn('Gestor de Plantilla', self.ventana('unidades').context['ver_mas'][0][0])

    def test_el_excel_saca_todas_las_filas_con_portada(self):
        from io import BytesIO

        from openpyxl import load_workbook
        respuesta = self.client.get(reverse('dashboard_ventana_excel', args=['activos']), {'sexo': 'F'})
        self.assertEqual(respuesta.status_code, 200)
        libro = load_workbook(BytesIO(respuesta.content))
        self.assertEqual(libro.sheetnames, ['Portada', 'Empleados'])
        hoja = libro['Empleados']
        self.assertEqual(hoja['B4'].value, 'Nombre')
        self.assertEqual(sorted(hoja.cell(row=f, column=2).value for f in (5, 6)),
                         ['Persona1 Prueba Test', 'Persona3 Prueba Test'])

    def test_las_ventanas_que_ya_existian_llevan_el_pie(self):
        for url, params in ((reverse('dashboard_salario'), {}),
                            (reverse('dashboard_altas_bajas_mes'), {'anno': 2026, 'mes': 6}),
                            (reverse('dashboard_piramide_detalle'), {'sexo': 'M', 'tramo': 0})):
            respuesta = self.client.get(url, params)
            self.assertContains(respuesta, 'data-excel-ventana')
        mes = self.client.get(reverse('dashboard_altas_bajas_mes'), {'anno': 2026, 'mes': 6})
        self.assertIn('alta_desde=2026-06-01&alta_hasta=2026-06-30', mes.context['ver_mas'][0][1])

