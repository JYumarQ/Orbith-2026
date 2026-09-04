# -*- coding: utf-8 -*-
"""Genera la plantilla .xlsx de «Altas y Bajas del Mes» (modelo de la UNE).

    .venv/Scripts/python.exe plantillas_excel/generadores/crear_altas_bajas.py

Escribe `plantillas_excel/informes_rrhh/altas_bajas_template.xlsx`, que es lo
que rellena `informes.excel.libro_de_altas_bajas`.

QUE LLEVA
---------
Una portada con el logo y los datos de identificacion, y las cuatro hojas del
modelo oficial en su orden y con sus cabeceras a dos alturas: Resumen Altas,
Altas, Resumen Bajas y Bajas.

Los nombres de las hojas se escriben aqui sin mes («Altas»). Al exportar se les
anade el del documento («Altas Julio 2026»), porque solo entonces se sabe cual
es. Excel no admite mas de 31 caracteres en el nombre de una hoja, y el mas
largo que se genera —«Resumen Bajas Septiembre 2026»— son 29.

LA FILA MODELO
--------------
Debajo de la cabecera de cada hoja de datos hay una fila que parece vacia y no
lo esta: lleva el formato con el que se escriben los datos. `informes/excel.py`
lo copia en cada fila que vuelca. Se puede cambiar su letra, su color o sus
bordes desde Excel —y los datos saldran asi—, pero si se BORRA la fila, las
filas de datos saldran sin formato.

EL RESTO SI SE EDITA A MANO
---------------------------
Titulos, rotulos de columna, anchos, colores y combinaciones estan en el fichero
y no en el codigo: se abre en Excel, se cambian y se guarda.
"""
import os
import sys

# El generador comparte con la exportacion la paleta y, sobre todo, en que fila
# empieza cada hoja. Tenerlo en un solo sitio es lo que impide que mover una
# cabecera aqui deje a la exportacion escribiendo encima de ella.
RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, RAIZ)
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'core.settings')

import django  # noqa: E402
django.setup()

from openpyxl import Workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from informes.excel import (CENTRADO, FILA_TITULO, FUENTE_CABECERA,  # noqa: E402
                            FUENTE_DATO, FUENTE_SUBCABECERA, FUENTE_TITULO,
                            HOJAS, IZQUIERDA, REJILLA, RELLENO_CABECERA,
                            RELLENO_SUBCABECERA)

DESTINO = os.path.join(RAIZ, 'plantillas_excel', 'informes_rrhh',
                       'altas_bajas_template.xlsx')


def cabecera(hoja, celda, texto, combinar=None, sub=False):
    """Un rotulo de cabecera, opcionalmente combinado con sus vecinas."""
    hoja[celda] = texto
    hoja[celda].font = FUENTE_SUBCABECERA if sub else FUENTE_CABECERA
    hoja[celda].fill = RELLENO_SUBCABECERA if sub else RELLENO_CABECERA
    hoja[celda].alignment = CENTRADO
    hoja[celda].border = REJILLA
    if combinar:
        hoja.merge_cells(combinar)
        # En un rango combinado, openpyxl solo pinta la celda superior izquierda:
        # el resto hay que bordearlas a mano o la rejilla se corta.
        for fila in hoja[combinar]:
            for celda_del_rango in fila:
                celda_del_rango.border = REJILLA
                celda_del_rango.fill = (RELLENO_SUBCABECERA if sub
                                        else RELLENO_CABECERA)


def titulo(hoja, celda, combinar):
    hoja[celda] = 'Altas y Bajas del Mes'
    hoja[celda].font = FUENTE_TITULO
    hoja[celda].alignment = IZQUIERDA
    hoja.merge_cells(combinar)


def fila_modelo(hoja, clave):
    """La fila cuyo formato heredan todas las de datos.

    El centrado se decide por columna, segun `centradas` de `HOJAS`: cifras,
    fechas y marcas de Si/No al centro; nombres, cargos y causas a la izquierda,
    que es como se lee un texto largo.
    """
    config = HOJAS[clave]
    for columna in range(1, config['columnas'] + 1):
        celda = hoja.cell(row=config['fila_datos'], column=columna)
        celda.font = FUENTE_DATO
        celda.border = REJILLA
        celda.alignment = (CENTRADO if columna in config['centradas']
                           else IZQUIERDA)


def anchos(hoja, medidas):
    for indice, medida in enumerate(medidas, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = medida


def hoja_de_portada(libro):
    """El contenido lo escribe la exportacion; aqui solo va el lienzo."""
    hoja = libro.create_sheet(HOJAS['portada']['nombre'])
    hoja.sheet_view.showGridLines = False
    hoja.row_dimensions[1].height = 58
    return hoja


def hoja_de_resumen(libro, clave, rotulo_total):
    """Municipio, los totales y los DOS desgloses, clave y no clave.

    El de puestos no clave repite exactamente la estructura del de puestos
    clave, a su derecha: mismo grupo de cabecera, misma pareja Cargo/Cantidad.
    La Provincia no esta: es la misma en todas las filas y la portada ya la
    declara una vez.
    """
    hoja = libro.create_sheet(HOJAS[clave]['nombre'])
    titulo(hoja, 'A%d' % FILA_TITULO, 'A%d:H%d' % (FILA_TITULO, FILA_TITULO))

    cabecera(hoja, 'A3', 'Municipio', 'A3:A5')
    cabecera(hoja, 'B3', rotulo_total, 'B3:B5')

    cabecera(hoja, 'C3', 'DE ELLAS EN CARGOS CLAVES', 'C3:E3')
    cabecera(hoja, 'C4', 'Total Puestos Clave', 'C4:C5')
    cabecera(hoja, 'D4', 'DESGLOSAR POR CARGOS', 'D4:E4')
    cabecera(hoja, 'D5', 'CARGO', sub=True)
    cabecera(hoja, 'E5', 'CANTIDAD DE CARGOS', sub=True)

    cabecera(hoja, 'F3', 'DE ELLAS EN PUESTOS NO CLAVE', 'F3:H3')
    cabecera(hoja, 'F4', 'Total Puestos NO Clave', 'F4:F5')
    cabecera(hoja, 'G4', 'DESGLOSAR POR CARGOS', 'G4:H4')
    cabecera(hoja, 'G5', 'CARGO', sub=True)
    cabecera(hoja, 'H5', 'CANTIDAD DE CARGOS', sub=True)

    anchos(hoja, [20, 14, 20, 52, 16, 22, 52, 16])
    hoja.row_dimensions[3].height = 24
    hoja.row_dimensions[4].height = 24
    hoja.row_dimensions[5].height = 30
    fila_modelo(hoja, clave)
    return hoja


def hoja_de_altas(libro):
    """13 columnas. Sin Provincia y con una sola columna de Si/No.

    El modelo oficial parte «Cargos Claves» en dos columnas, «Si» y «No», y se
    marca con una X la que toca. Aqui va en una sola que dice Si o No: ocupa la
    mitad y se lee sin buscar donde cayo la X.
    """
    hoja = libro.create_sheet(HOJAS['altas']['nombre'])
    titulo(hoja, 'B%d' % FILA_TITULO, 'B%d:M%d' % (FILA_TITULO, FILA_TITULO))
    hoja['B3'] = 'Empresa:'
    hoja['B3'].font = FUENTE_SUBCABECERA

    columnas = [
        ('A', 'No'), ('B', 'Nombre y Apellidos'), ('C', 'Cargo'),
        ('D', 'Area en la que Laborara'),
        ('E', 'Cargos Claves (SPxR diferenciado)'),
        ('F', 'Sexo'), ('G', 'Edad'), ('H', 'Cat. Ocup.'),
        ('I', 'Nivel escolar'), ('J', 'Graduado en:'), ('K', 'Fecha del alta'),
        ('L', 'Direccion / Area o UEB'), ('M', 'Municipio'),
    ]
    for letra, texto in columnas:
        cabecera(hoja, '%s4' % letra, texto, '%s4:%s5' % (letra, letra))

    # «Cargos Claves (SPxR diferenciado)» necesita ancho de sobra: es el rotulo
    # mas largo de la hoja y ya no lo reparte entre dos columnas.
    anchos(hoja, [5, 34, 40, 30, 18, 7, 7, 11, 12, 26, 14, 26, 18])
    hoja.row_dimensions[4].height = 44
    fila_modelo(hoja, 'altas')
    return hoja


def hoja_de_bajas(libro):
    """18 columnas. Sin Provincia y con las tres marcas en una sola columna."""
    hoja = libro.create_sheet(HOJAS['bajas']['nombre'])
    titulo(hoja, 'B%d' % FILA_TITULO, 'B%d:R%d' % (FILA_TITULO, FILA_TITULO))
    hoja['B4'] = 'EMPRESA'
    hoja['B4'].font = FUENTE_SUBCABECERA

    columnas = [
        ('A', 'No'), ('B', 'Nombre y Apellidos'), ('C', 'Cargo'),
        ('D', 'Cargos Claves (SPxR diferenciado)'),
        ('E', 'Actividad que realizaba'), ('F', 'Area en la que Laboraba'),
        ('G', 'Sexo'), ('H', 'Edad'), ('I', 'Cat. Ocup.'),
        ('J', 'Nivel escolar'), ('K', 'Graduado en:'),
        ('L', 'Fecha de la baja'), ('M', 'Causa de la Baja'),
        ('N', 'Marcar si requiere completar el cargo'),
        ('O', 'Anos en la Organizacion'), ('P', 'Direccion / Area o UEB'),
        ('Q', 'Municipio'),
        ('R', 'Sistema de pago por resultados (cobrado por el trabajador)'),
    ]
    for letra, texto in columnas:
        cabecera(hoja, '%s5' % letra, texto, '%s5:%s6' % (letra, letra))

    # Las tres columnas de rotulo largo —D, N y R— llevan ancho de sobra: antes
    # repartian el titulo entre dos columnas y ahora tiene que caber en una.
    anchos(hoja, [5, 34, 40, 18, 30, 30, 7, 7, 11, 12, 26, 14, 26,
                  20, 22, 26, 18, 24])
    hoja.row_dimensions[5].height = 58
    fila_modelo(hoja, 'bajas')
    return hoja


libro = Workbook()
libro.remove(libro.active)          # la hoja vacia que openpyxl crea sola

hoja_de_portada(libro)
hoja_de_resumen(libro, 'resumen_altas', 'TOTAL')
hoja_de_altas(libro)
hoja_de_resumen(libro, 'resumen_bajas', 'TOTAL BAJAS')
hoja_de_bajas(libro)

# Todas apaisadas: son tablas anchas y en vertical no cabe ni la mitad. Y con la
# cabecera inmovilizada, para que al bajar por una lista de cuarenta personas se
# siga sabiendo que columna se esta mirando.
POR_NOMBRE = {config['nombre']: config for config in HOJAS.values()}
for hoja in libro.worksheets:
    hoja.page_setup.orientation = 'landscape'
    hoja.page_setup.fitToWidth = 1
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    fila_datos = POR_NOMBRE[hoja.title]['fila_datos']
    if fila_datos:
        hoja.freeze_panes = hoja.cell(row=fila_datos, column=1)

os.makedirs(os.path.dirname(DESTINO), exist_ok=True)
libro.save(DESTINO)
print('plantilla escrita en', DESTINO)
print('hojas:', libro.sheetnames)
