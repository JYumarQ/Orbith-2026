# -*- coding: utf-8 -*-
"""Genera la plantilla .xlsx de los Anexos 17 y 17.1.

    .venv/Scripts/python.exe plantillas_excel/generadores/crear_anexo.py

Escribe `plantillas_excel/informes_rrhh/anexo_template.xlsx`, que es lo que
rellena `informes_anexo.excel.libro`.

QUÉ LLEVA: TRES HOJAS MODELO
----------------------------
El fichero no es el libro final: es el MOLDE del que sale. Trae tres hojas y
ninguna se entrega tal cual:

  * `Portada`     — la rellena `informes.excel._escribir_portada()`, la misma de
                    los siete informes de RRHH. Aquí solo se reserva la hoja.
  * `Modelo 17`   — el detalle: una fila por cargo, seis columnas.
  * `Modelo 17.1`  — el conteo: una fila por unidad o departamento, siete columnas.

Al exportar se copia el molde que toque una vez por unidad, se renombra y se
rellena; al final los moldes se retiran del libro. Copiarlos en vez de dibujar
cada hoja por código es lo que hace que todas salgan idénticas y que retocar el
molde en Excel las cambie todas a la vez.

LAS TRES FILAS MODELO DE CADA HOJA DE DATOS
-------------------------------------------
Debajo de la cabecera hay tres filas que parecen vacías y no lo están: llevan el
formato de los tres tipos de fila del informe —el detalle, el rótulo de grupo y
el subtotal—. La exportación copia su formato en cada fila que vuelca. Se puede
cambiar su letra, su color o sus bordes desde Excel y los datos saldrán así; si
se BORRAN, saldrán sin formato.

EL RESTO SÍ SE EDITA A MANO
---------------------------
Títulos, rótulos, anchos, colores y combinaciones están en el fichero y no en el
código: se abre en Excel, se cambian y se guarda. Lo único que no se puede mover
son las FILAS donde empieza cada cosa, porque el código escribe en ellas: están
declaradas en `informes_anexo/excel/libro.py` y se importan aquí para que no
puedan discrepar.
"""
import os
import sys

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, RAIZ)
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'core.settings')

import django  # noqa: E402
django.setup()

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from informes.excel import (AZUL, CENTRADO, FILA_ALTERNA,  # noqa: E402
                            FUENTE_CABECERA, FUENTE_DATO, FUENTE_ROTULO,
                            FUENTE_TITULO, IZQUIERDA, REJILLA,
                            RELLENO_CABECERA, RELLENO_GRUPO,
                            RELLENO_SUBCABECERA)

from informes_anexo.excel.libro import (DISENOS, FILA_CABECERA,  # noqa: E402
                                        FILA_EMPRESA, FILA_MODELO,
                                        FILA_ORGANISMO, FILA_UNIDAD,
                                        HOJA_PORTADA, NOMBRE_PLANTILLA)
from informes_anexo.anexo17_1 import (NIVEL_BLOQUE, NIVEL_DETALLE,  # noqa: E402
                                     NIVEL_SUBTOTAL)

DESTINO = os.path.join(RAIZ, 'plantillas_excel', 'informes_rrhh',
                       NOMBRE_PLANTILLA)

# Sangría de la columna de denominación en cada tipo de fila. El detalle va
# metido hacia dentro respecto al rótulo que lo agrupa: es lo que deja ver la
# jerarquía de un vistazo, sin líneas ni símbolos.
SANGRIA = {NIVEL_DETALLE: 2, NIVEL_BLOQUE: 0, NIVEL_SUBTOTAL: 1}


def titulo(hoja, texto, columnas):
    """Las dos líneas de título, centradas sobre el ancho de la tabla."""
    for fila, valor, fuente in ((1, texto, FUENTE_TITULO),
                                (2, 'Plantilla de Cargos', FUENTE_ROTULO)):
        hoja.merge_cells(start_row=fila, start_column=1,
                         end_row=fila, end_column=columnas)
        celda = hoja.cell(row=fila, column=1)
        celda.value = valor
        celda.font = fuente
        celda.alignment = CENTRADO


def identificacion(hoja, columnas, columna_derecha):
    """Las tres filas institucionales, vacías y ya formateadas.

    El contenido lo pone la exportación desde Parámetros Generales. Escribirlo
    aquí ataría la plantilla a una entidad concreta.
    """
    for fila in (FILA_EMPRESA, FILA_ORGANISMO, FILA_UNIDAD):
        hoja.merge_cells(start_row=fila, start_column=1,
                         end_row=fila, end_column=columna_derecha - 1)
        hoja.merge_cells(start_row=fila, start_column=columna_derecha,
                         end_row=fila, end_column=columnas)
        for columna in (1, columna_derecha):
            celda = hoja.cell(row=fila, column=columna)
            celda.font = FUENTE_DATO
            celda.alignment = IZQUIERDA


def cabecera(hoja, rotulos):
    """Los rótulos de columna."""
    for indice, texto in enumerate(rotulos, start=1):
        celda = hoja.cell(row=FILA_CABECERA, column=indice)
        celda.value = texto
        celda.font = FUENTE_CABECERA
        celda.fill = RELLENO_CABECERA
        celda.alignment = CENTRADO
        celda.border = REJILLA
    hoja.row_dimensions[FILA_CABECERA].height = 28


def filas_modelo(hoja, columnas, centradas):
    """Las tres filas que llevan el formato de los datos."""
    formatos = {
        NIVEL_DETALLE: (FUENTE_DATO, PatternFill('solid', fgColor=FILA_ALTERNA)),
        NIVEL_BLOQUE: (Font(name='Calibri', size=10, bold=True, color=AZUL),
                       RELLENO_GRUPO),
        NIVEL_SUBTOTAL: (Font(name='Calibri', size=10, bold=True),
                         RELLENO_SUBCABECERA),
    }

    for nivel, fila in FILA_MODELO.items():
        fuente, relleno = formatos[nivel]
        for columna in range(1, columnas + 1):
            celda = hoja.cell(row=fila, column=columna)
            celda.font = fuente
            celda.fill = relleno
            celda.border = REJILLA
            if columna == 1:
                celda.alignment = Alignment(horizontal='left', vertical='center',
                                            indent=SANGRIA[nivel])
            elif columna in centradas:
                celda.alignment = CENTRADO
            else:
                celda.alignment = IZQUIERDA


def impresion(hoja, orientacion):
    """A un ancho de página y repitiendo la cabecera en cada hoja impresa."""
    hoja.page_setup.orientation = orientacion
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.print_title_rows = '%d:%d' % (FILA_CABECERA, FILA_CABECERA)
    hoja.freeze_panes = hoja.cell(row=FILA_MODELO[NIVEL_DETALLE], column=1)


libro = Workbook()

# La portada va primera y la rellena `informes.excel._escribir_portada()`, que
# escribe en B2, B3 y de la fila 6 hacia abajo. Aquí solo se reserva la hoja: su
# aspecto lo decide esa función, compartida con los siete informes de RRHH.
portada = libro.active
portada.title = HOJA_PORTADA
portada.sheet_view.showGridLines = False

for diseno in DISENOS.values():
    hoja = libro.create_sheet(title=diseno['molde'])
    columnas = len(diseno['rotulos'])
    titulo(hoja, diseno['titulo'], columnas)
    identificacion(hoja, columnas, diseno['columna_derecha'])
    cabecera(hoja, diseno['rotulos'])
    filas_modelo(hoja, columnas, diseno['centradas'])
    impresion(hoja, diseno['orientacion'])
    for indice, ancho in enumerate(diseno['anchos'], start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho

libro.save(DESTINO)
print('Plantilla escrita en %s' % DESTINO)
