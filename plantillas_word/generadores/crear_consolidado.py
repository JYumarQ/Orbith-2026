# -*- coding: utf-8 -*-
"""Genera la plantilla .docx del «Consolidado de la Fuerza de Trabajo».

    .venv/Scripts/python.exe plantillas_word/generadores/crear_consolidado.py

Escribe `plantillas_word/informes_economia/consolidado_template.docx`, que es lo
que rellena `informes.views.ExportarConsolidadoWordView`.

POR QUE SE REHIZO
-----------------
Era la unica plantilla de los siete informes que nadie habia fabricado: venia de
un .docx escrito a mano, sin generador, y su exportacion se armaba un contexto
propio sin la cabecera comun. El documento salia sin logo, sin provincia y sin
decir quien lo habia exportado.

QUE LLEVA Y EN QUE ORDEN
------------------------
Lo mismo que la pantalla, el PDF y el Excel: las cifras del periodo, la tabla de
las cinco categorias ocupacionales con sus once columnas, y el resumen final.

LA TABLA NO TIENE BUCLE
-----------------------
Las cinco categorias son FIJAS —Obreros, Tecnico, Cuadro, Servicio y
Administrativo— y no salen de un recorrido: son las que reconoce
`InformeEconomiaView.mapear_categoria`. Por eso esta plantilla no lleva ninguna
etiqueta `{%tr for %}` y se puede retocar entera en Word sin miedo a borrar una
fila de control.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm  # noqa: E402

from comun import (AZUL, BLANCO, GRIS, NARANJA, PIZARRA, VERDE,  # noqa: E402
                   FONDO_CABECERA_OSCURA, cabecera_de_tabla,
                   documento_apaisado, escribir, lista_de_unidades, sombrear,
                   tarjetas, titulo)

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DESTINO = os.path.join(RAIZ, 'informes_economia', 'consolidado_template.docx')

# Las cinco filas fijas de la tabla. El nombre es la clave con la que la vista
# las guarda en `datos`; tiene que coincidir con `mapear_categoria`.
CATEGORIAS = ('Obreros', 'Técnico', 'Cuadro', 'Servicio', 'Administrativo')

# Las once columnas: rotulo, ancho y color de la cifra.
COLUMNAS = (
    ('Cargo u Ocupación', Cm(4.4), PIZARRA),
    ('Plantilla Aprobada', Cm(2.2), PIZARRA),
    ('Plantilla Cubierta', Cm(2.2), VERDE),
    ('Plantilla Vacante', Cm(2.2), NARANJA),
    ('Mujeres', Cm(1.9), AZUL),
    ('Determinados', Cm(2.1), PIZARRA),
    ('Determinados Mujeres', Cm(2.1), AZUL),
    ('Adiestrados', Cm(2.0), PIZARRA),
    ('Adiestrados Mujeres', Cm(2.0), AZUL),
    ('Total', Cm(2.0), PIZARRA),
    ('Total Mujeres', Cm(1.9), NARANJA),
)

# Las claves de cada cifra, en el mismo orden que las columnas de datos.
CLAVES = ('aprob', 'cub', 'vac', 'p_muj', 'd_tot', 'd_muj',
          'a_tot', 'a_muj', 't_tot', 't_muj')


documento = documento_apaisado()
titulo(documento, 'Consolidado de la Fuerza de Trabajo')
lista_de_unidades(documento)

escribir(documento.add_paragraph(), 'Período: {{ periodo_texto }}', 12,
         negrita=True, color=GRIS)

tarjetas(documento, [
    ('Trabajadores', '{{ totales.t_tot }}', PIZARRA),
    ('Mujeres', '{{ totales.t_muj }}', AZUL),
    ('Plantilla aprobada', '{{ totales.aprob }}', VERDE),
    ('Vacantes', '{{ totales.vac }}', NARANJA),
])

escribir(documento.add_paragraph(),
         'Plantilla y contratos por categoría ocupacional', 13, negrita=True,
         color=AZUL, mayusculas=True)

#   0  cabecera de columnas    1..5  una fila por categoria    6  totales
tabla = documento.add_table(rows=2 + len(CATEGORIAS), cols=len(COLUMNAS))
tabla.style = 'Table Grid'
tabla.autofit = False
cabecera_de_tabla(tabla,
                  [(rotulo, indice == 0)
                   for indice, (rotulo, _, _) in enumerate(COLUMNAS)],
                  [ancho for _, ancho, _ in COLUMNAS])

for numero, categoria in enumerate(CATEGORIAS, start=1):
    fila = tabla.rows[numero]
    escribir(fila.cells[0].paragraphs[0], categoria, 11, negrita=True,
             color=PIZARRA)
    for indice, clave in enumerate(CLAVES, start=1):
        parrafo = fila.cells[indice].paragraphs[0]
        parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        escribir(parrafo, '{{ datos.%s.%s }}' % (categoria, clave), 11,
                 negrita=True, color=COLUMNAS[indice][2])

total = tabla.rows[-1]
escribir(total.cells[0].paragraphs[0], 'Total', 12, negrita=True, color=BLANCO,
         mayusculas=True)
for indice, clave in enumerate(CLAVES, start=1):
    parrafo = total.cells[indice].paragraphs[0]
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    escribir(parrafo, '{{ totales.%s }}' % clave, 12, negrita=True,
             color=BLANCO)
for celda in total.cells:
    sombrear(celda, FONDO_CABECERA_OSCURA)

documento.add_paragraph()
escribir(documento.add_paragraph(), 'Resumen', 13, negrita=True, color=AZUL,
         mayusculas=True)
tarjetas(documento, [
    ('Total de trabajadores', '{{ totales.t_tot }}', PIZARRA),
    ('De ellos, mujeres', '{{ totales.t_muj }}', AZUL),
    ('Plazas vacantes', '{{ totales.vac }}', NARANJA),
])

os.makedirs(os.path.dirname(DESTINO), exist_ok=True)
documento.save(DESTINO)
print('plantilla escrita en', DESTINO)
print('columnas:', len(COLUMNAS), '| filas de la tabla:', len(tabla.rows))
