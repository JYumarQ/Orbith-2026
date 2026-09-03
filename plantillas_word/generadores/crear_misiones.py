# -*- coding: utf-8 -*-
"""Genera la plantilla .docx de «Misiones Internacionales».

    .venv/Scripts/python.exe plantillas_word/generadores/crear_misiones.py

Escribe `plantillas_word/informes_economia/misiones_template.docx`, que es lo
que rellena `informes.views.ExportarMisionesWordView`.

EL .docx SE PUEDE EDITAR A MANO
-------------------------------
Es un documento de Word normal: fuentes, colores y anchos se cambian con las
herramientas de siempre. Lo unico intocable son las cuatro filas de la tabla que
llevan etiquetas `{%tr ... %}`: parecen vacias y son las que repiten los datos.
El mapa de filas esta abajo, junto al codigo que las escribe.

Si se retoca el .docx a mano, conviene traer el mismo cambio aqui; si no, el
siguiente que ejecute este generador deshara el retoque sin enterarse.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm  # noqa: E402

from comun import (AZUL, INDIGO, NARANJA, PIZARRA, VERDE, FONDO_GRUPO,  # noqa: E402
                   cabecera_de_tabla, documento_apaisado, escribir,
                   fila_de_control, lista_de_unidades, sombrear, tarjetas,
                   titulo)

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DESTINO = os.path.join(RAIZ, 'informes_economia', 'misiones_template.docx')

# Anchos: el nombre y el cargo se llevan el sitio, que es donde estan los textos
# largos; el resto son codigos cortos.
ANCHOS = [Cm(3.0), Cm(7.0), Cm(8.0), Cm(4.0), Cm(1.6), Cm(2.3)]
COLUMNAS = [('Exp. Laboral', False), ('Nombre y Apellidos', True),
            ('Cargo u Ocupación', True), ('País', False), ('Cat.', False),
            ('G. Escala', False)]

documento = documento_apaisado()
titulo(documento, 'Personal en Misión Internacional')
lista_de_unidades(documento)

tarjetas(documento, [
    ('Total en Misión', '{{ total_misiones }}', AZUL),
    ('Países', '{{ paises }}', VERDE),
    ('Unidades', '{{ unidades_con_mision }}', NARANJA),
    ('Importe Total', '{{ "%.2f"|format(importe_total) }}', INDIGO),
])

#   0  cabecera de columnas (se repite en cada pagina)
#   1  for unidad          2  nombre de la unidad
#   3  for trabajador      4  datos del trabajador     5  endfor
#   6  endfor
tabla = documento.add_table(rows=7, cols=len(COLUMNAS))
tabla.style = 'Table Grid'
tabla.autofit = False
cabecera_de_tabla(tabla, COLUMNAS, ANCHOS)

# Fila de unidad: ocupa toda la anchura, como en la pantalla y el PDF.
fila_de_control(tabla, 1, '{%tr for unidad in misiones_por_unidad %}')
fila_unidad = tabla.rows[2]
celda_unidad = fila_unidad.cells[0].merge(fila_unidad.cells[len(COLUMNAS) - 1])
escribir(celda_unidad.paragraphs[0],
         '{{ unidad.nombre }} — {{ unidad.total }} en misión',
         12, negrita=True, color=INDIGO)
sombrear(celda_unidad, FONDO_GRUPO)

fila_de_control(tabla, 3, '{%tr for t in unidad.trabajadores %}')
fila_datos = tabla.rows[4]
valores = [('{{ t.expediente }}', False, PIZARRA),
           ('{{ t.nombre }}', True, PIZARRA),
           ('{{ t.cargo }}', True, PIZARRA),
           ('{{ t.pais }}', False, AZUL),
           ('{{ t.cat_letra }}', False, PIZARRA),
           ('{{ t.g_escala }}', False, PIZARRA)]
for i, (valor, a_la_izquierda, color) in enumerate(valores):
    parrafo = fila_datos.cells[i].paragraphs[0]
    parrafo.alignment = (WD_ALIGN_PARAGRAPH.LEFT if a_la_izquierda
                         else WD_ALIGN_PARAGRAPH.CENTER)
    escribir(parrafo, valor, 11, color=color)

fila_de_control(tabla, 5, '{%tr endfor %}')
fila_de_control(tabla, 6, '{%tr endfor %}')

os.makedirs(os.path.dirname(DESTINO), exist_ok=True)
documento.save(DESTINO)
print('plantilla escrita en', DESTINO)
print('filas de la tabla:', len(tabla.rows))
