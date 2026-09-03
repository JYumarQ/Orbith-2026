# -*- coding: utf-8 -*-
"""Genera la plantilla .docx de «Registro Diario por Trabajadores».

    .venv/Scripts/python.exe plantillas_word/generadores/crear_registro_diario.py

Escribe `plantillas_word/informes_economia/registro_diario_template.docx`, que
es lo que rellena `informes.views.ExportarRegistroDiarioWordView`.

QUE LLEVA Y QUE NO
------------------
* La Eficiencia del Plan va ARRIBA y a todo lo ancho, no a la derecha como en la
  pantalla: en papel el ancho es fijo, y una columna lateral deja la tabla
  estrecha y la eficiencia perdida en un lado.
* Los graficos NO van: son de JavaScript y no existen fuera del navegador. Lo
  que se archiva son las cifras.

EL .docx SE PUEDE EDITAR A MANO
-------------------------------
La tabla de indicadores lleva UNA sola pareja de etiquetas `{%tr for %}` /
`{%tr endfor %}`, asi que se abre en Word y se retoca sin riesgo. Lo unico que
no se toca son esas dos filas, que parecen vacias.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm  # noqa: E402

from comun import (AZUL, BLANCO, GRIS, NARANJA, PIZARRA, VERDE,  # noqa: E402
                   cabecera_de_tabla, documento_apaisado, escribir,
                   fila_de_control, lista_de_unidades, sombrear, titulo)

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DESTINO = os.path.join(RAIZ, 'informes_economia',
                       'registro_diario_template.docx')

AZUL_PLAN = '002B6B'

documento = documento_apaisado()
titulo(documento, 'Registro Diario por Trabajadores')
lista_de_unidades(documento)

# --- Eficiencia del Plan, arriba y a todo lo ancho --------------------------
escribir(documento.add_paragraph(), 'Eficiencia del Plan', 13, negrita=True,
         color=AZUL, mayusculas=True)

CIFRAS = [('Plan', '{{ plan_editable }}', BLANCO),
          ('Real', '{{ plan_real }}', BLANCO),
          ('Cumplimiento', '{{ plan_cumplimiento }}%', BLANCO),
          ('Diferencia', '{{ plan_diferencia }}', BLANCO),
          ('Mujeres', '{{ total_mujeres }}', BLANCO)]

plan = documento.add_table(rows=2, cols=len(CIFRAS))
for columna, (rotulo, valor, color) in enumerate(CIFRAS):
    arriba = plan.cell(0, columna).paragraphs[0]
    arriba.alignment = WD_ALIGN_PARAGRAPH.CENTER
    escribir(arriba, rotulo, 10, negrita=True, color=BLANCO, mayusculas=True)
    abajo = plan.cell(1, columna).paragraphs[0]
    abajo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    escribir(abajo, valor, 24, negrita=True, color=color)
    sombrear(plan.cell(0, columna), AZUL_PLAN)
    sombrear(plan.cell(1, columna), AZUL_PLAN)

documento.add_paragraph()

# --- Indicadores oficiales -------------------------------------------------
escribir(documento.add_paragraph(), 'Indicadores oficiales', 12, negrita=True,
         color=AZUL, mayusculas=True)

COLUMNAS = [('Indicadores de Control', True), ('Trabajadores', False),
            ('De ellos, Mujeres', False)]
ANCHOS = [Cm(14.0), Cm(5.0), Cm(5.0)]

#   0  cabecera de columnas (se repite en cada pagina)
#   1  for indicador     2  datos del indicador     3  endfor
#   4  fila de totales
tabla = documento.add_table(rows=5, cols=len(COLUMNAS))
tabla.style = 'Table Grid'
tabla.autofit = False
cabecera_de_tabla(tabla, COLUMNAS, ANCHOS)

fila_de_control(tabla, 1, '{%tr for nombre, cifras in datos.items() %}')
fila = tabla.rows[2]
escribir(fila.cells[0].paragraphs[0], '{{ nombre }}', 11, negrita=True,
         color=PIZARRA)
for i, (valor, color) in enumerate((('{{ cifras.tot }}', VERDE),
                                    ('{{ cifras.muj }}', AZUL)), start=1):
    parrafo = fila.cells[i].paragraphs[0]
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    escribir(parrafo, valor, 11, negrita=True, color=color)
fila_de_control(tabla, 3, '{%tr endfor %}')

total = tabla.rows[4]
escribir(total.cells[0].paragraphs[0], 'Total de trabajadores físicos', 12,
         negrita=True, color=BLANCO, mayusculas=True)
for i, valor in enumerate(('{{ plan_real }}', '{{ total_mujeres }}'), start=1):
    parrafo = total.cells[i].paragraphs[0]
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    escribir(parrafo, valor, 12, negrita=True, color=BLANCO)
for celda in total.cells:
    sombrear(celda, '0F172A')

os.makedirs(os.path.dirname(DESTINO), exist_ok=True)
documento.save(DESTINO)
print('plantilla escrita en', DESTINO)
print('filas de la tabla:', len(tabla.rows))
