"""Importar y exportar el nomenclador de cargos en Excel (admin de Django).

Formato (el mismo de «Cargos OK.xlsx»): hoja «CARGOS_CODIGO» (o la primera) con
las columnas Cargos · Grupo Escala · Categoría Ocupacional · Salario Básico.
Reglas comunes en `core/excel_admin.py`: todo o nada, simular, nunca borra.

Mejoras sobre el importador anterior (24-10-2026):
* revisa el Excel ENTERO antes de guardar: con un error no se guarda nada;
* un cargo se reconoce por su NOMBRE y su GRUPO DE ESCALA: en la empresa hay
  cargos que se llaman igual y tienen otro grupo, y son cargos distintos (antes se
  reconocía solo por el nombre y el segundo se perdía en silencio);
* avisa de las filas REPETIDAS (mismo nombre y mismo grupo), diciendo si traen
  datos distintos;
* dice qué cambiaría en cada cargo que ya existe, y permite simular;
* exporta en el mismo formato, para descargar, corregir y volver a subir;
* ya no necesita pandas para leer.
"""

from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import load_workbook

from core.excel_admin import Simulacion, clave, columnas, hoja_con_titulos, libro_vacio, limpio

from .models import NCargo, NGrupoEscala
from .utils import obtener_categoria_codigo

HOJA = 'CARGOS_CODIGO'
COLUMNAS = [
    ('cargo', ('Cargos', 'Cargo')),
    ('grupo', ('Grupo Escala', 'Grupo de Escala', 'Grupo')),
    ('categoria', ('Categoría Ocupacional', 'Categoria', 'Categoría')),
    ('salario', ('Salario Básico', 'Salario Basico', 'Salario')),
]


def _categoria(texto):
    """Código de la categoría a partir del nombre («Técnico») o del propio código («TEC»)."""
    codigos = dict(NCargo._meta.get_field('cat_ocupacional').choices)
    if texto.upper() in codigos:
        return texto.upper()
    return obtener_categoria_codigo(texto)


def _leer(archivo):
    """`(filas, repetidos, errores)`. `filas`: {(nombre, grupo): (fila, datos)} con la
    primera aparición de cada cargo."""
    libro = load_workbook(archivo, read_only=True, data_only=True)
    hoja = libro[HOJA] if HOJA in libro.sheetnames else libro.worksheets[0]
    fila_titulos, col = columnas(hoja, COLUMNAS)
    grupos = {clave(g.nivel): g for g in NGrupoEscala.objects.all()}

    filas, repetidos, errores = {}, [], []
    for numero, valores in enumerate(hoja.iter_rows(min_row=fila_titulos + 1, values_only=True),
                                     start=fila_titulos + 1):
        def celda(nombre):
            i = col[nombre]
            return limpio(valores[i]) if i < len(valores) else ''

        nombre = celda('cargo')
        if not nombre:
            continue
        grupo = grupos.get(clave(celda('grupo')))
        categoria = _categoria(celda('categoria'))
        try:
            salario = Decimal(celda('salario').replace('$', '').replace(',', '') or '0').quantize(Decimal('0.01'))
        except InvalidOperation:
            salario = None
        problemas = []
        if not celda('grupo'):
            problemas.append('falta el grupo de escala')
        elif grupo is None:
            problemas.append('el grupo de escala «%s» no existe' % celda('grupo'))
        if not celda('categoria'):
            problemas.append('falta la categoría ocupacional')
        elif categoria is None:
            problemas.append('la categoría «%s» no se reconoce' % celda('categoria'))
        if salario is None or salario < 0:
            problemas.append('el salario «%s» no es una cifra válida' % celda('salario'))
        if problemas:
            errores.append('Fila %d («%s»): %s.' % (numero, nombre, '; '.join(problemas)))
            continue

        datos = {'descripcion': nombre, 'grupo_escala': grupo, 'cat_ocupacional': categoria,
                 'salario_basico': salario}
        k = (clave(nombre), clave(grupo.nivel))
        if k in filas:
            primera, anteriores = filas[k]
            iguales = all(anteriores[c] == datos[c] for c in ('cat_ocupacional', 'salario_basico'))
            repetidos.append('Fila %d: «%s» (grupo %s) ya estaba en la fila %d%s. Se usa la fila %d.' % (
                numero, nombre, grupo.nivel, primera,
                '' if iguales else ' con OTROS datos (categoría %s, salario %s; aquí categoría %s, salario %s)' % (
                    anteriores['cat_ocupacional'], anteriores['salario_basico'], categoria, salario),
                primera))
            continue
        filas[k] = (numero, datos)
    return filas, repetidos, errores


def _diferencias(cargo, datos):
    cambios = []
    if cargo.grupo_escala_id != datos['grupo_escala'].pk:
        cambios.append('grupo %s → %s' % (cargo.grupo_escala.nivel if cargo.grupo_escala_id else '—',
                                           datos['grupo_escala'].nivel))
    if cargo.cat_ocupacional != datos['cat_ocupacional']:
        cambios.append('categoría %s → %s' % (cargo.cat_ocupacional, datos['cat_ocupacional']))
    if (cargo.salario_basico or Decimal('0')) != datos['salario_basico']:
        cambios.append('salario %s → %s' % (cargo.salario_basico, datos['salario_basico']))
    return cambios


def importar_cargos(archivo, estrategia='SALTAR', simular=False):
    """Importa el Excel. `estrategia`: 'SALTAR' (los cargos que ya existen no se
    tocan) o 'ACTUALIZAR' (se ponen como dice el Excel). Devuelve el resultado
    para la pantalla; con errores no guarda nada."""
    resultado = {'creados': [], 'actualizados': [], 'saltados': [], 'sin_cambios': 0,
                 'repetidos': [], 'errores': [], 'simulado': simular, 'guardado': False}
    try:
        filas, resultado['repetidos'], resultado['errores'] = _leer(archivo)
    except (ValueError, KeyError, OSError) as error:
        resultado['errores'].append(str(error))
        return resultado
    except Exception:
        resultado['errores'].append('No se pudo leer el fichero: ¿es un Excel (.xlsx)?')
        return resultado
    if resultado['errores']:
        return resultado

    # Mismo nombre y mismo grupo = el mismo cargo. Si hubiera dos iguales en la base,
    # se toma el más antiguo (el otro ya era un duplicado antes de importar).
    existentes = {}
    for c in NCargo.objects.select_related('grupo_escala').order_by('-pk'):
        existentes[(clave(c.descripcion), clave(c.grupo_escala.nivel if c.grupo_escala_id else ''))] = c
    try:
        with transaction.atomic():
            for k, (numero, datos) in filas.items():
                cargo = existentes.get(k)
                if cargo is None:
                    NCargo.objects.create(**datos)
                    resultado['creados'].append(datos['descripcion'])
                    continue
                cambios = _diferencias(cargo, datos)
                if not cambios:
                    resultado['sin_cambios'] += 1
                elif estrategia != 'ACTUALIZAR':
                    resultado['saltados'].append('%s: %s' % (cargo.descripcion, ', '.join(cambios)))
                else:
                    for campo, valor in datos.items():
                        setattr(cargo, campo, valor)
                    cargo.save()
                    resultado['actualizados'].append('%s: %s' % (cargo.descripcion, ', '.join(cambios)))
            if simular:
                raise Simulacion()
        resultado['guardado'] = True
    except Simulacion:
        pass
    return resultado


def exportar_cargos():
    """El nomenclador en el mismo formato que lee `importar_cargos`."""
    etiquetas = dict(NCargo._meta.get_field('cat_ocupacional').choices)
    filas = [(c.descripcion, c.grupo_escala.nivel if c.grupo_escala_id else '',
              etiquetas.get(c.cat_ocupacional, c.cat_ocupacional), float(c.salario_basico or 0))
             for c in NCargo.objects.select_related('grupo_escala').order_by('cat_ocupacional', 'descripcion')]
    libro = libro_vacio()
    hoja_con_titulos(libro, HOJA, ['Cargos', 'Grupo Escala', 'Categoría Ocupacional', 'Salario Básico'],
                     filas, [60, 14, 24, 16])
    return libro
