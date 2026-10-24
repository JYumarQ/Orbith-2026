# -*- coding: utf-8 -*-
"""Los libros de Excel de los Anexos 17 y 17.1.

TRES LIBROS, UN SOLO MOLDE
--------------------------
Se exportan tres cosas distintas y las tres salen de aquí:

  * `MODO_17`    — el Anexo 17 solo, el detalle.
  * `MODO_17_1`   — el Anexo 17.1 solo, el conteo.
  * `MODO_AMBOS` — los dos juntos, como el fichero que entregó Capital Humano.

Los tres empiezan igual: **Portada** y **Resumen**. Detrás van las hojas de cada
unidad —`Oficina Central 17`, `Oficina Central 17.1`, `Florida 17`…— según el modo.

POR QUÉ NO SE DIBUJA LA HOJA
----------------------------
Se abre una PLANTILLA `.xlsx` y se rellena, igual que hace
`informes.excel.libro_de_altas_bajas()` con el modelo de la UNE. Son modelos
oficiales que Capital Humano tiene que poder retocar —un ancho, un color, un
rótulo— sin pedir un cambio de código, y reproducir a mano sus cabeceras
combinadas desde Python daría un resultado parecido al papel, no idéntico.

La plantilla la fabrica `plantillas_excel/generadores/crear_anexo.py`, que
importa de aquí las posiciones: si estuvieran declaradas en los dos sitios,
mover una cabecera en el generador dejaría a este módulo escribiendo encima.

QUÉ ENTRA EN EL LIBRO
---------------------
Lo que el usuario haya elegido en el selector de unidades. Sin nada elegido,
todo su alcance —que es el libro completo de siempre—. Las unidades SIN plantilla
no generan hoja: una pestaña vacía en un documento que se archiva no dice «esta
unidad no tiene plazas», dice «aquí falta algo».
"""
from copy import copy

from django.http import HttpResponse
from openpyxl import load_workbook
from openpyxl.styles import Font

from informes.excel import (FUENTE_DATO, FUENTE_PIE, IZQUIERDA,
                            LARGO_MAXIMO_HOJA, TIPO_MIME, _escribir_portada,
                            ruta_de_plantilla)

from ..anexo17_1 import (COLUMNAS_17_1, NIVEL_BLOQUE, NIVEL_DETALLE,
                        NIVEL_SUBTOTAL)

NOMBRE_PLANTILLA = 'anexo_template.xlsx'
HOJA_PORTADA = 'Portada'
HOJA_RESUMEN = 'Resumen'

MODO_17 = '17'
MODO_17_1 = '17_1'
MODO_17_2 = '17_2'
MODO_AMBOS = 'ambos'

#: Cómo se llama cada informe en el nombre del fichero y en el de la hoja.
NOMBRES_DE_INFORME = {
    MODO_17: 'Anexo 17',
    MODO_17_1: 'Anexo 17.1',
    MODO_17_2: 'Anexo 17.2',
    MODO_AMBOS: 'Anexo 17 y 17.1',
}

# --------------------------------------------------------------------------
# Dónde está cada cosa en la plantilla
# --------------------------------------------------------------------------
# Compartido con el generador. Son las mismas filas en las dos hojas de datos:
# los dos modelos oficiales tienen el mismo encabezado y solo cambian de columnas.

FILA_EMPRESA = 3
FILA_ORGANISMO = 4
FILA_UNIDAD = 5
FILA_CABECERA = 7

#: Las tres filas modelo. La primera es también donde empiezan los datos.
FILA_MODELO = {
    NIVEL_DETALLE: 8,
    NIVEL_BLOQUE: 9,
    NIVEL_SUBTOTAL: 10,
}
FILA_DATOS = FILA_MODELO[NIVEL_DETALLE]

#: Los rótulos de las columnas de conteo salen de la MISMA definición que usa la
#: pantalla, y no de una lista escrita aquí a mano: duplicarlos era la forma
#: segura de que un día el papel y la pantalla dejaran de llamar igual a lo mismo.
ROTULOS_17_1 = (['Denominación']
               + [rotulo for _, rotulo, _, _ in COLUMNAS_17_1]
               + ['Total de Cargos'])

ROTULOS_17 = ['Cargos actual', 'Categ. ocup.', 'Cantidad de cargos', 'Rol',
              'Nivel de preparación', 'Grupo Escala']

DISENOS = {
    MODO_17: {
        'molde': 'Modelo 17',
        'sufijo': '17',
        'titulo': 'Anexo No 17',
        'rotulos': ROTULOS_17,
        # Ancho en caracteres. La primera carga con nombres de cargo largos.
        'anchos': [52, 12, 14, 14, 26, 12],
        'centradas': (2, 3, 4, 6),
        'columna_derecha': 5,
        'orientacion': 'landscape',
    },
    MODO_17_1: {
        'molde': 'Modelo 17.1',
        'sufijo': '17.1',
        'titulo': 'Anexo No 17.1',
        'rotulos': ROTULOS_17_1,
        'anchos': [46, 11, 8, 8, 8, 8, 16],
        'centradas': (2, 3, 4, 5, 6, 7),
        'columna_derecha': 5,
        'orientacion': 'portrait',
    },
}


# --------------------------------------------------------------------------
# Utilidades de hoja
# --------------------------------------------------------------------------

def _estilos_modelo(hoja, columnas):
    """El formato de las tres filas modelo, clonado para poder repartirlo.

    Se copia con `copy()` y no se asigna tal cual: en openpyxl el `_style` de una
    celda es un objeto compartido, y reutilizarlo haría que cambiar el formato de
    una fila cambiara el de todas las demás.
    """
    return {
        nivel: [copy(hoja.cell(row=fila, column=col)._style)
                for col in range(1, columnas + 1)]
        for nivel, fila in FILA_MODELO.items()
    }


def _nombre_de_hoja(nombre, usados):
    """Un nombre de pestaña válido y único.

    Excel rechaza los de más de 31 caracteres y los que llevan `/ \\ ? * : [ ]`.
    El nombre de la unidad va ENTERO y solo se recorta si no cabe: abreviarlo
    siempre —«Oficina C 17»— ahorraba cuatro letras a costa de que dos unidades
    de nombre parecido se confundieran.

    Y rechaza dos hojas con el mismo nombre: dos unidades pueden coincidir en los
    primeros 31 caracteres, así que a la segunda se le añade un sufijo en vez de
    dejar que reviente el guardado.
    """
    limpio = str(nombre or 'Hoja')
    for prohibido in '/\\?*:[]':
        limpio = limpio.replace(prohibido, '-')
    limpio = limpio.strip()[:LARGO_MAXIMO_HOJA] or 'Hoja'

    candidato, sufijo = limpio, 2
    while candidato.lower() in usados:
        cola = ' (%d)' % sufijo
        candidato = limpio[:LARGO_MAXIMO_HOJA - len(cola)] + cola
        sufijo += 1

    usados.add(candidato.lower())
    return candidato


def _escribir_cabecera(hoja, datos, unidad, columnas, columna_derecha):
    """Las seis líneas de identificación de una hoja.

    Empresa, organismo, provincia y REUP son los mismos en todo el libro y salen
    de Parámetros Generales. La unidad y el municipio son de ESTA hoja: el
    Resumen habla de la empresa y cada hoja de unidad habla de la suya.
    """
    hoja.cell(row=FILA_EMPRESA, column=1,
              value='Empresa: %s' % (datos.get('cab_empresa') or '—'))
    hoja.cell(row=FILA_EMPRESA, column=columna_derecha,
              value='Código REUP: %s' % (datos.get('cab_reup') or '—'))

    hoja.cell(row=FILA_ORGANISMO, column=1,
              value='Organismo: %s' % (datos.get('cab_organismo') or '—'))
    hoja.cell(row=FILA_ORGANISMO, column=columna_derecha,
              value='Provincia: %s' % (datos.get('cab_provincia') or '—'))

    nombre = unidad.descripcion if unidad else (datos.get('cab_empresa') or '—')
    hoja.cell(row=FILA_UNIDAD, column=1, value='Unidad: %s' % nombre)

    # El municipio solo existe cuando la hoja va de UNA unidad. En el Resumen se
    # escribe la marca de hueco en vez de dejar la celda vacía: en un documento
    # que se archiva, un blanco no distingue «no aplica» de «se olvidó».
    municipio = unidad.municipio.nombre if unidad and unidad.municipio_id else ''
    hoja.cell(row=FILA_UNIDAD, column=columna_derecha,
              value='Municipio: %s' % (municipio or '—'))


def _volcar(hoja, fila_inicial, estilos, filas, columnas):
    """Escribe las filas con el formato que le toca a cada nivel.

    Cada fila es `(nivel, [valores])`. Devuelve la siguiente fila libre.
    """
    numero = fila_inicial
    for nivel, valores in filas:
        formato = estilos.get(nivel, estilos[NIVEL_DETALLE])
        for col in range(1, columnas + 1):
            celda = hoja.cell(row=numero, column=col)
            celda._style = copy(formato[col - 1])
            celda.value = valores[col - 1] if col <= len(valores) else None
        numero += 1
    return numero


def _rotulo_de_seccion(hoja, fila, texto, columnas):
    """Una banda a todo lo ancho, para separar bloques dentro de la hoja."""
    hoja.merge_cells(start_row=fila, start_column=1,
                     end_row=fila, end_column=columnas)
    celda = hoja.cell(row=fila, column=1)
    celda.value = texto
    celda._style = copy(hoja.cell(row=FILA_MODELO[NIVEL_BLOQUE], column=1)._style)
    celda.alignment = IZQUIERDA
    return fila + 1


def _escribir_firmas(hoja, fila_inicial, datos, columnas):
    """El pie de conformidad: quién la propone y quién la aprueba."""
    fecha = datos.get('fecha_generacion')
    fecha_texto = fecha.strftime('%d/%m/%Y') if fecha else ''

    lineas = [
        ('Propuesto por: %s' % _nombre_firmante(datos.get('director_ch')),
         _titulo_firmante(datos.get('director_ch'), 'Director/a de Capital Humano')),
        ('Aprobado por: %s' % _nombre_firmante(datos.get('director_gral')),
         _titulo_firmante(datos.get('director_gral'), 'Director/a General')),
    ]

    numero = fila_inicial + 2
    for encabezado, cargo in lineas:
        hoja.merge_cells(start_row=numero, start_column=1,
                         end_row=numero, end_column=columnas)
        celda = hoja.cell(row=numero, column=1)
        celda.value = '%s        Firma: _______________        Fecha: %s' % (
            encabezado, fecha_texto)
        celda.font = Font(name='Calibri', size=10, bold=True)
        celda.alignment = IZQUIERDA
        numero += 1

        hoja.merge_cells(start_row=numero, start_column=1,
                         end_row=numero, end_column=columnas)
        celda = hoja.cell(row=numero, column=1)
        celda.value = cargo
        celda.font = FUENTE_DATO
        celda.alignment = IZQUIERDA
        numero += 2

    hoja.merge_cells(start_row=numero, start_column=1,
                     end_row=numero, end_column=columnas)
    celda = hoja.cell(row=numero, column=1)
    celda.value = 'Generado por Orbith — %s' % fecha_texto
    celda.font = FUENTE_PIE
    celda.alignment = IZQUIERDA


def _nombre_firmante(firmante):
    """El nombre completo, o un hueco marcado si no se resolvió."""
    aspirante = (firmante or {}).get('aspirante')
    if not aspirante:
        return '_______________________'
    partes = [aspirante.nombre, aspirante.papellido, aspirante.sapellido or '']
    return ' '.join(parte for parte in partes if parte).strip()


def _titulo_firmante(firmante, defecto):
    return (firmante or {}).get('titulo') or defecto


# --------------------------------------------------------------------------
# Las filas de cada modelo
# --------------------------------------------------------------------------

def _filas_de_conteo(bloque_filas):
    """Las filas del 17.1, ya como `(nivel, valores)`.

    Las de rótulo van sin cifras: sus números son los del `Subtotal` que cierra
    su bloque, y repetirlos arriba haría que la columna pareciera sumar el doble.
    """
    salida = []
    for fila in bloque_filas:
        if fila['nivel'] == NIVEL_BLOQUE:
            valores = [fila['denominacion']] + [None] * len(COLUMNAS_17_1) + [None]
        else:
            valores = ([fila['denominacion']]
                       + [fila[clave] for clave, _, _, _ in COLUMNAS_17_1]
                       + [fila['total']])
        salida.append((fila['nivel'], valores))
    return salida


def _filas_de_detalle(secciones, etiquetas):
    """Las filas del Anexo 17: rótulos de subunidad y departamento, y sus cargos.

    Devuelve también el total de plazas, que es lo que cierra la hoja. Los cargos
    fuera de plantilla NO entran aquí: van en su propio bloque, después del Total.
    """
    filas, total = [], 0

    for seccion in secciones:
        if not seccion["unidad"].__dict__.get("es_la_propia_unidad"):
            filas.append((NIVEL_BLOQUE, [seccion["unidad"].descripcion]))

        for depto in seccion["departamentos"]:
            filas.append((NIVEL_SUBTOTAL, [depto.descripcion]))
            for cargo in depto.cargos_filtrados:
                cantidad = cargo.cant_aprobada or 0
                total += cantidad
                ncargo = cargo.ncargo
                escala = getattr(ncargo.grupo_escala, 'nivel', '') if ncargo.grupo_escala_id else ''
                filas.append((NIVEL_DETALLE, [
                    ncargo.descripcion,
                    etiquetas.get(ncargo.cat_ocupacional, ncargo.cat_ocupacional),
                    cantidad,
                    cargo.rol.tipo if cargo.rol_id else '—',
                    cargo.nivel_preparacion.nombre if cargo.nivel_preparacion_id else '',
                    escala,
                ]))

    return filas, total


def _filas_fuera_de_detalle(secciones, etiquetas):
    """Lo mismo, pero con los cargos que quedaron fuera de la plantilla."""
    filas, total = [], 0

    for seccion in secciones:
        deptos = seccion["departamentos_fuera"]
        if not deptos:
            continue
        if not seccion["unidad"].__dict__.get("es_la_propia_unidad"):
            filas.append((NIVEL_BLOQUE, [seccion["unidad"].descripcion]))
        for depto in deptos:
            filas.append((NIVEL_SUBTOTAL, [depto.descripcion]))
            for cargo in depto.cargos_filtrados:
                cantidad = cargo.cant_aprobada or 0
                total += cantidad
                ncargo = cargo.ncargo
                escala = getattr(ncargo.grupo_escala, 'nivel', '') if ncargo.grupo_escala_id else ''
                filas.append((NIVEL_DETALLE, [
                    ncargo.descripcion,
                    etiquetas.get(ncargo.cat_ocupacional, ncargo.cat_ocupacional),
                    cantidad,
                    cargo.rol.tipo if cargo.rol_id else '—',
                    cargo.nivel_preparacion.nombre if cargo.nivel_preparacion_id else '',
                    escala,
                ]))

    return filas, total


# --------------------------------------------------------------------------
# El libro
# --------------------------------------------------------------------------

def _rellenar(hoja, datos, unidad, diseno, filas, total_fila, fuera, rotulo_fuera):
    """Una hoja completa: cabecera, tabla, bloque de fuera de plantilla y firmas.

    El bloque de fuera de plantilla va AL FINAL, después del Total, y con su
    propio total: esas plazas existen y pueden tener gente, pero no forman parte
    de la plantilla aprobada por el organismo y no suman en la cifra oficial.
    """
    columnas = len(diseno['rotulos'])
    _escribir_cabecera(hoja, datos, unidad, columnas, diseno['columna_derecha'])

    # Los formatos se leen ANTES de escribir: las filas modelo están justo donde
    # van a caer los primeros datos y se sobrescriben enseguida.
    estilos = _estilos_modelo(hoja, columnas)

    numero = _volcar(hoja, FILA_DATOS, estilos, filas, columnas)
    numero = _volcar(hoja, numero, estilos, [total_fila], columnas)

    filas_fuera, total_fuera = fuera
    if filas_fuera:
        numero = _rotulo_de_seccion(hoja, numero + 1, rotulo_fuera, columnas)
        numero = _volcar(hoja, numero, estilos, filas_fuera, columnas)
        numero = _volcar(hoja, numero, estilos, [total_fuera], columnas)

    _escribir_firmas(hoja, numero, datos, columnas)


def _ajustes_de_impresion(hoja, molde):
    """Repone en la hoja copiada lo que `copy_worksheet` deja atrás.

    Los ajustes de impresión y la fila de cabecera repetida viven fuera de la
    rejilla de celdas, así que no se copian: sin esto, cada hoja saldría con los
    de fábrica y la cabecera no se repetiría al imprimir.
    """
    hoja.page_setup.orientation = molde.page_setup.orientation
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.print_title_rows = '%d:%d' % (FILA_CABECERA, FILA_CABECERA)
    hoja.freeze_panes = hoja.cell(row=FILA_DATOS, column=1)


def libro_del_anexo(datos, modo):
    """El libro completo del modo pedido.

    `datos['unidades']` es una lista de dicts, uno por unidad, con lo que hace
    falta para pintar sus dos hojas. Lo arma `views._datos_del_libro()`, que es
    quien sabe consultar; aquí solo se escribe.
    """
    libro = load_workbook(ruta_de_plantilla(NOMBRE_PLANTILLA))
    moldes = {clave: libro[diseno['molde']] for clave, diseno in DISENOS.items()}
    usados = {HOJA_PORTADA.lower()}

    # --- Portada. La misma de los siete informes de RRHH, para que un documento
    # exportado de Orbith se reconozca como tal venga del módulo que venga. ---
    # `fecha_generacion` viaja como `date` porque las firmas y el nombre del
    # fichero la formatean cada uno a su manera. La portada la ESCRIBE en una
    # celda, y una fecha suelta ahí sale como «2026-09-15 00:00:00»: se le pasa
    # ya en texto, igual que hacen los siete informes de RRHH.
    contexto_portada = dict(datos, fecha_generacion=datos.get('fecha_texto') or '')
    _escribir_portada(
        libro[HOJA_PORTADA], contexto_portada,
        datos.get('fecha_texto') or '',
        titulo=NOMBRES_DE_INFORME[modo],
        cifras=datos.get('cifras_portada') or [])

    # --- Resumen. Lo llevan los TRES libros: el conteo por unidad es la puerta
    # de entrada al documento, se pida el detalle o se pida el resumen. ---
    diseno_resumen = DISENOS[MODO_17_1]
    hoja = libro.copy_worksheet(moldes[MODO_17_1])
    hoja.title = _nombre_de_hoja(HOJA_RESUMEN, usados)
    _ajustes_de_impresion(hoja, moldes[MODO_17_1])
    resumen = datos['resumen']
    resumen_fuera = datos['resumen_fuera']
    _rellenar(
        hoja, datos, None, diseno_resumen,
        _filas_de_conteo(resumen['filas']),
        _filas_de_conteo([resumen['total']])[0],
        (_filas_de_conteo(resumen_fuera['filas']) if resumen_fuera['total']['total'] else [],
         _filas_de_conteo([resumen_fuera['total']])[0]),
        'CARGOS FUERA DE LA PLANTILLA APROBADA (no incluidos en el total anterior)')

    # --- Una hoja por unidad y modo. ---
    modos = [MODO_17, MODO_17_1] if modo == MODO_AMBOS else [modo]

    for entrada in datos['unidades']:
        unidad = entrada['unidad']
        for clave in modos:
            diseno = DISENOS[clave]
            hoja = libro.copy_worksheet(moldes[clave])
            hoja.title = _nombre_de_hoja(
                '%s %s' % (unidad.descripcion, diseno['sufijo']), usados)
            _ajustes_de_impresion(hoja, moldes[clave])

            if clave == MODO_17_1:
                aprobada, fuera = entrada['conteo'], entrada['conteo_fuera']
                _rellenar(
                    hoja, datos, unidad, diseno,
                    _filas_de_conteo(aprobada['filas']),
                    _filas_de_conteo([aprobada['total']])[0],
                    (_filas_de_conteo(fuera['filas']) if fuera['total']['total'] else [],
                     _filas_de_conteo([fuera['total']])[0]),
                    'CARGOS FUERA DE LA PLANTILLA APROBADA '
                    '(no incluidos en el total anterior)')
            else:
                etiquetas = datos.get('etiquetas_categoria') or {}
                filas, total = _filas_de_detalle(entrada['secciones'], etiquetas)
                filas_fuera, total_fuera = _filas_fuera_de_detalle(
                    entrada['secciones'], etiquetas)
                _rellenar(
                    hoja, datos, unidad, diseno, filas,
                    (NIVEL_SUBTOTAL, ['Total', None, total]),
                    (filas_fuera, (NIVEL_SUBTOTAL, ['Total fuera de plantilla',
                                                    None, total_fuera])),
                    'CARGOS FUERA DE LA PLANTILLA APROBADA '
                    '(no incluidos en el Total anterior)')

    for molde in moldes.values():
        libro.remove(molde)

    return libro


def nombre_de_fichero(datos, modo, filtrado=False):
    """`Anexo_17_Toda_la_Empresa_15-09-2026.xlsx`, con el alcance de verdad.

    Primero el informe, después el alcance y la fecha, como el resto de
    descargas de Orbith (core/nombres_descarga.py, 24-10-2026).

    El alcance no puede ser el nombre de «la» unidad desde que el selector admite
    varias: con dos marcadas, cualquiera de los dos nombres mentiría. Se dice lo
    que hay —«Toda la Empresa», el nombre si es una sola, «3 unidades» si son
    varias— y así el fichero archivado se explica solo.

    `filtrado` lo marca en el nombre porque un libro con filtros NO es la
    plantilla de la unidad: es un recorte, y confundirlos en un documento que se
    entrega al organismo sería grave.
    """
    from core.nombres_descarga import nombre_de_informe
    partes = [NOMBRES_DE_INFORME[modo], datos.get('alcance_fichero') or 'Informe']
    if filtrado:
        partes.append('filtrado')
    partes.append((datos.get('fecha_generacion').strftime('%d-%m-%Y')
                   if datos.get('fecha_generacion') else ''))
    return nombre_de_informe(*partes, extension='xlsx', fecha=False)


def respuesta_excel(libro, nombre_fichero):
    """El libro, ya como descarga del navegador."""
    salida = HttpResponse(content_type=TIPO_MIME)
    salida['Content-Disposition'] = 'attachment; filename="%s"' % nombre_fichero
    libro.save(salida)
    return salida
