# -*- coding: utf-8 -*-
"""El libro de Excel del Anexo 17.2.

POR QUÉ NO COMPARTE PLANTILLA CON EL 17 Y EL 17.1
-------------------------------------------------
Se decidió que el 17.2 va en libro APARTE, y además su hoja tiene otra forma: 17
columnas en vez de 6 o 7, la cabecera a dos alturas —«Salario» abarcando siete
columnas y sus subtítulos debajo—, una fila de numeración, y DOS filas por plaza
en vez de una. Forzar todo eso dentro del molde de los otros dos habría llenado
`libro.py` de excepciones para el caso raro.

Lo que sí se comparte es todo lo demás: el nombre de hoja, la portada y la
respuesta HTTP se importan de `libro.py`, que las tiene resueltas.

LAS DOS FILAS POR PLAZA
-----------------------
Debajo del importe mensual va la tarifa horaria, como en el fichero de Capital
Humano: el importe entre el fondo de tiempo de Parámetros Generales. `CAlta` no
guarda la tarifa, se deriva. El cálculo se hace aquí y NO se importa de
Pre-Nómina: los Anexos y ese módulo son independientes y no se tocan.

Solo se dividen las columnas que llevan un importe del mes. El CLA y la
nocturnidad ya vienen en CUP/hora desde las tarjetas del contrato, así que su
casilla de la segunda fila se queda vacía: dividir una tarifa horaria entre las
horas del mes no significa nada.
"""
from copy import copy

from openpyxl import load_workbook

from informes.excel import FUENTE_PIE, IZQUIERDA, _escribir_portada, ruta_de_plantilla

from .libro import HOJA_PORTADA, _nombre_de_hoja, respuesta_excel  # noqa: F401

NOMBRE_PLANTILLA = 'anexo_17_2_template.xlsx'
MOLDE = 'Modelo 17.2'

TITULO = ('ANEXO No.17.2 De la plantilla de cargos '
          'y el registro de trabajadores')

#: Los 17 rótulos, repartidos en las dos alturas de la cabecera. Los siete del
#: medio cuelgan de «Salario», que los abarca en la fila de arriba.
ROTULOS_FIJOS = ['No.', 'Descripción órganos y cargos', 'Categoría ocupacional',
                 'Nombre y apellidos\ndel trabajador', 'Sexo',
                 'Carné de Identidad', 'Nivel de preparación', 'Grupo escala']
ROTULOS_SALARIO = ['Total', 'Escala', 'CLA', 'T. noct y mixtos',
                   'Maestría\no doctorado', 'Años de servicio',
                   'Otros (CNCI, Pago Instructores)']
ROTULOS_COLA = ['Rol', 'Tridente']

COLUMNAS = len(ROTULOS_FIJOS) + len(ROTULOS_SALARIO) + len(ROTULOS_COLA)  # 17
PRIMERA_SALARIO = len(ROTULOS_FIJOS) + 1   # columna I
ULTIMA_SALARIO = PRIMERA_SALARIO + len(ROTULOS_SALARIO) - 1  # columna O

# Dónde está cada cosa. Compartido con
# `plantillas_excel/generadores/crear_anexo_17_2.py`.
FILA_TITULO = 1
FILA_ENTIDAD = 2
FILA_FECHA = 3
FILA_CABECERA_ALTA = 4
FILA_CABECERA_BAJA = 5
FILA_NUMEROS = 6
FILA_DATOS = 7
#: La segunda fila modelo: la de la tarifa horaria que va bajo cada plaza.
FILA_MODELO_TARIFA = 8

#: Formato del modelo que entregó RRHH el 24-10-2026: casi todas las cabeceras
#: van giradas 90° (ver `ROTULOS_HORIZONTALES`), así que las columnas pueden ser
#: estrechas y la hoja cabe en el ancho de la página.
ANCHOS = [7.4, 28.3, 6.6, 29, 5.4, 11.7, 24.1, 5.9,
          7.6, 8.9, 7.1, 8.1, 10.1, 8.7, 9.3,
          3.4, 3.9]

#: Las únicas cabeceras que NO van giradas: las de texto largo y «Salario», que
#: abarca siete columnas.
ROTULOS_HORIZONTALES = {'Descripción órganos y cargos',
                        'Nombre y apellidos\ndel trabajador', 'Salario'}
#: Alto de las dos filas de la cabecera: la de abajo, alta para el texto girado.
ALTO_CABECERA_ALTA = 31.5
ALTO_CABECERA_BAJA = 122.25
#: Tamaño de letra de toda la hoja, como en el modelo.
TAMANO_LETRA = 12


def _estilos(hoja):
    """El formato de las dos filas modelo, clonado para poder repartirlo."""
    return {
        'dato': [copy(hoja.cell(row=FILA_DATOS, column=c)._style)
                 for c in range(1, COLUMNAS + 1)],
        'tarifa': [copy(hoja.cell(row=FILA_MODELO_TARIFA, column=c)._style)
                   for c in range(1, COLUMNAS + 1)],
    }


def _valores(fila):
    """Una plaza, ya como lista de 17 celdas.

    Una vacante escribe el cargo, la categoría, el grupo escala y el rol, y deja
    en blanco lo que depende de la persona. El salario va en blanco y no en cero:
    un cero es una cifra y aquí no hay ninguna que dar.
    """
    vacia = fila['vacante']
    return [
        fila['expediente'] or None,
        fila['cargo'],
        fila['categoria'],
        fila['nombre'] or None,
        fila['sexo'] or None,
        fila['carne'] or None,
        fila['formacion'] or None,
        fila['grupo_escala'],
        None if vacia else fila['total'],
        None if vacia else fila['escala'],
        # Tarifa por hora de la tarjeta CLA, y las dos de nocturnidad en
        # texto separadas por una barra. No son importes del mes.
        None if vacia else (fila['cla'] or None),
        None if vacia else (fila['nocturnos'] or None),
        None if vacia else fila['maestria_doctorado'],
        None if vacia else fila['annos_servicio'],
        None if vacia else fila['otros'],
        # El modelo oficial escribe la INICIAL del rol («F» de Fundamental,
        # «A» de Apoyo), no la palabra entera: la columna es estrecha y el
        # organismo lo lee así. En pantalla, donde hay sitio, sí va completa.
        (fila['rol'][:1].upper() if fila['rol'] else None),
        fila['tridente'] or None,
    ]


def _fila_de_tarifa(fila, fondo):
    """La segunda fila: el importe dividido entre el fondo de tiempo.

    Solo llevan cifra las columnas de dinero; el resto va en blanco para que la
    fila se lea como lo que es, una continuación de la de arriba y no una plaza
    nueva. Una vacante no la lleva: no hay importe que dividir.
    """
    if fila['vacante'] or not fondo:
        return None

    valores = [None] * COLUMNAS
    # `None` en las dos casillas que no llevan un importe del mes: el CLA y la
    # nocturnidad ya son CUP/hora y no hay nada que dividir.
    claves = ['total', 'escala', None, None, 'maestria_doctorado',
              'annos_servicio', 'otros']
    for desplazamiento, clave in enumerate(claves):
        if clave is None:
            continue
        importe = fila[clave]
        if importe:
            valores[PRIMERA_SALARIO - 1 + desplazamiento] = float(importe) / float(fondo)
    return valores


def _escribir_cabecera(hoja, datos, hoja_datos):
    """Las tres líneas de identificación de la hoja."""
    hoja.cell(row=FILA_ENTIDAD, column=1,
              value='Entidad: %s' % (datos.get('cab_empresa') or '—'))

    depto = hoja_datos.get('departamento')
    if depto is not None:
        alcance = '%s (%s)' % (depto.unidad_organizativa.descripcion,
                               depto.descripcion)
    else:
        # Las tres hojas especiales no son un departamento: se nombran con la
        # unidad y lo que agrupan, en el mismo formato que las demás.
        alcance = '%s (%s)' % (hoja_datos['unidad'].descripcion,
                               hoja_datos['nombre'])
    hoja.cell(row=FILA_ENTIDAD, column=4,
              value='Establecimiento, Taller o Unidad Administrativa: %s' % alcance)

    hoja.cell(row=FILA_FECHA, column=4,
              value='Fecha de aprobación: %s' % (datos.get('fecha_texto') or ''))


def _rellenar(hoja, datos, hoja_datos, fondo):
    estilos = _estilos(hoja)
    _escribir_cabecera(hoja, datos, hoja_datos)

    numero = FILA_DATOS
    for fila in hoja_datos['filas']:
        valores = _valores(fila)
        for col in range(1, COLUMNAS + 1):
            celda = hoja.cell(row=numero, column=col)
            celda._style = copy(estilos['dato'][col - 1])
            celda.value = valores[col - 1]
        numero += 1

        tarifa = _fila_de_tarifa(fila, fondo)
        if tarifa is not None:
            for col in range(1, COLUMNAS + 1):
                celda = hoja.cell(row=numero, column=col)
                celda._style = copy(estilos['tarifa'][col - 1])
                celda.value = tarifa[col - 1]
            numero += 1

    numero += 1
    hoja.merge_cells(start_row=numero, start_column=1,
                     end_row=numero, end_column=COLUMNAS)
    celda = hoja.cell(row=numero, column=1)
    celda.value = 'Generado por Orbith — %s' % (datos.get('fecha_texto') or '')
    celda.font = FUENTE_PIE
    celda.alignment = IZQUIERDA


def libro_del_anexo_17_2(datos):
    """El libro completo: portada y una hoja por departamento de cada unidad.

    Sin hoja de Resumen: ese conteo es el del 17.1 y en un registro de
    trabajadores no dice gran cosa.
    """
    libro = load_workbook(ruta_de_plantilla(NOMBRE_PLANTILLA))
    molde = libro[MOLDE]
    usados = {HOJA_PORTADA.lower()}
    fondo = datos.get('fondo_tiempo')

    contexto_portada = dict(datos, fecha_generacion=datos.get('fecha_texto') or '')
    _escribir_portada(libro[HOJA_PORTADA], contexto_portada,
                      datos.get('fecha_texto') or '',
                      titulo='Anexo 17.2',
                      cifras=datos.get('cifras_portada') or [])

    for hoja_datos in datos['hojas']:
        hoja = libro.copy_worksheet(molde)
        # Con varias unidades a la vez, dos pueden tener un departamento que se
        # llame igual; `_nombre_de_hoja` le pone sufijo a la segunda. La cabecera
        # de dentro dice «Unidad (Departamento)», así que no hay ambigüedad.
        hoja.title = _nombre_de_hoja(hoja_datos['nombre'], usados)
        hoja.page_setup.orientation = molde.page_setup.orientation
        hoja.page_setup.fitToWidth = 1
        hoja.page_setup.fitToHeight = 0
        hoja.sheet_properties.pageSetUpPr.fitToPage = True
        hoja.print_title_rows = '%d:%d' % (FILA_CABECERA_ALTA, FILA_NUMEROS)
        hoja.freeze_panes = hoja.cell(row=FILA_DATOS, column=1)
        _rellenar(hoja, datos, hoja_datos, fondo)

    libro.remove(molde)
    return libro
