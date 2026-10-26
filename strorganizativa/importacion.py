"""Importar y exportar la estructura organizativa en Excel (admin de Django, 24-10-2026).

Un libro con dos hojas:

* **Unidades**: Código · Unidad · Tipo · Pertenece a · Municipio · Grupo de nómina ·
  Posición. Obligatorias: Unidad y Tipo (y Grupo de nómina en una unidad nueva que
  no sea subunidad; «Pertenece a» en una subunidad).
* **Departamentos**: Unidad · Departamento · Posición. Obligatorias las dos primeras.

Una unidad se reconoce por su Código si lo trae, y si no por su nombre (sin tener
en cuenta tildes ni mayúsculas). Con el Código se puede RENOMBRAR una unidad; sin él,
un nombre que no existe es una unidad nueva. Un departamento se reconoce por su
nombre dentro de su unidad. La Posición vacía pone lo nuevo al final.

Reglas comunes en `core/excel_admin.py`: todo o nada, simular, nunca borra, y lo
que exporta se puede volver a importar tal cual.
"""

from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.db.models import Max
from openpyxl import load_workbook

from core.excel_admin import Simulacion, clave, columnas, hoja_con_titulos, libro_vacio, limpio

from .models import (MENSAJE_UNIDAD_PRINCIPAL, Departamento, UnidadOrganizativa,
                     unidad_con_el_grupo_de_nomina)

HOJA_UNIDADES = 'Unidades'
HOJA_DEPARTAMENTOS = 'Departamentos'
COL_UNIDADES = [('unidad', ('Unidad', 'Unidad Organizativa', 'Descripción')), ('tipo', ('Tipo', 'Tipo de unidad'))]
COL_UNIDADES_OPC = [('codigo', ('Código', 'Codigo interno', 'Código interno')),
                    ('padre', ('Pertenece a', 'Unidad principal', 'Padre')),
                    ('municipio', ('Municipio',)), ('grupo', ('Grupo de nómina', 'Grupo nomina')),
                    ('posicion', ('Posición', 'Orden'))]
COL_DPTOS = [('unidad', ('Unidad', 'Unidad Organizativa')), ('departamento', ('Departamento',))]
COL_DPTOS_OPC = [('posicion', ('Posición', 'Orden'))]


def _entero(texto):
    texto = limpio(texto)
    if not texto:
        return None
    try:
        numero = float(texto)
    except ValueError:
        raise ValueError(texto)
    if numero != int(numero) or numero < 0:
        raise ValueError(texto)
    return int(numero)


def _filas(hoja, requeridas, opcionales):
    fila_titulos, col = columnas(hoja, requeridas, opcionales)
    for numero, valores in enumerate(hoja.iter_rows(min_row=fila_titulos + 1, values_only=True),
                                     start=fila_titulos + 1):
        celdas = {n: (valores[i] if i < len(valores) else None) for n, i in col.items()}
        if any(limpio(v) for v in celdas.values()):
            yield numero, celdas


def _leer(archivo):
    from configuracion.models import Configuracion
    from nomencladores.models import NMunicipio, NTipoUnidadOrganizativa

    libro = load_workbook(archivo, read_only=True, data_only=True)
    if HOJA_UNIDADES not in libro.sheetnames:
        raise ValueError('El libro debe tener una hoja llamada «%s» (y, si quiere cargar departamentos, '
                         'otra llamada «%s»). Descargue el Excel con «Exportar Excel» para ver el formato.'
                         % (HOJA_UNIDADES, HOJA_DEPARTAMENTOS))

    tipos = {clave(t.descripcion): t for t in NTipoUnidadOrganizativa.objects.all()}
    config = Configuracion.objects.first()
    municipios = NMunicipio.objects.all()
    if config and config.provincia_entidad_id:
        municipios = municipios.filter(provincia_id=config.provincia_entidad_id)
    municipios = {clave(m.nombre): m for m in municipios}

    unidades, errores = [], []
    nombres = {}
    for numero, c in _filas(libro[HOJA_UNIDADES], COL_UNIDADES, COL_UNIDADES_OPC):
        nombre = limpio(c.get('unidad'))
        problemas = []
        tipo = tipos.get(clave(c.get('tipo')))
        if not nombre:
            problemas.append('falta el nombre de la unidad')
        if not limpio(c.get('tipo')):
            problemas.append('falta el tipo')
        elif tipo is None:
            problemas.append('el tipo «%s» no existe' % limpio(c.get('tipo')))
        municipio = None
        if limpio(c.get('municipio')):
            municipio = municipios.get(clave(c.get('municipio')))
            if municipio is None:
                problemas.append('el municipio «%s» no es de la provincia de la entidad' % limpio(c.get('municipio')))
        datos = {'fila': numero, 'nombre': nombre, 'tipo': tipo, 'municipio': municipio,
                 'padre': limpio(c.get('padre'))}
        for campo, etiqueta in (('codigo', 'código'), ('grupo', 'grupo de nómina'), ('posicion', 'posición')):
            try:
                datos[campo] = _entero(c.get(campo))
            except ValueError as error:
                problemas.append('el %s «%s» no es un número entero' % (etiqueta, error))
        if tipo is not None and tipo.es_subunidad and not datos['padre']:
            problemas.append('es una subunidad: diga en «Pertenece a» de qué unidad principal')
        if nombre and clave(nombre) in nombres:
            problemas.append('la unidad está repetida (ya está en la fila %d)' % nombres[clave(nombre)])
        if problemas:
            errores.append('Unidades, fila %d («%s»): %s.' % (numero, nombre or '—', '; '.join(problemas)))
            continue
        nombres[clave(nombre)] = numero
        unidades.append(datos)

    departamentos, vistos = [], {}
    if HOJA_DEPARTAMENTOS in libro.sheetnames:
        for numero, c in _filas(libro[HOJA_DEPARTAMENTOS], COL_DPTOS, COL_DPTOS_OPC):
            unidad, nombre = limpio(c.get('unidad')), limpio(c.get('departamento'))
            problemas = []
            if not unidad:
                problemas.append('falta la unidad')
            if not nombre:
                problemas.append('falta el nombre del departamento')
            try:
                posicion = _entero(c.get('posicion'))
            except ValueError as error:
                posicion = None
                problemas.append('la posición «%s» no es un número entero' % error)
            k = (clave(unidad), clave(nombre))
            if nombre and k in vistos:
                problemas.append('está repetido en la misma unidad (fila %d)' % vistos[k])
            if problemas:
                errores.append('Departamentos, fila %d («%s»): %s.' % (numero, nombre or '—', '; '.join(problemas)))
                continue
            vistos[k] = numero
            departamentos.append({'fila': numero, 'unidad': unidad, 'nombre': nombre, 'posicion': posicion})
    return unidades, departamentos, errores


ETIQUETAS = {'descripcion': 'nombre', 'tipo': 'tipo', 'padre': 'pertenece a', 'municipio': 'municipio',
             'grupo_nomina': 'grupo de nómina', 'orden_informe': 'posición'}


def _siguiente(queryset):
    return (queryset.aggregate(m=Max('orden_informe'))['m'] or 0) + 1


def importar_estructura(archivo, estrategia='SALTAR', simular=False):
    """Importa unidades y departamentos. `estrategia`: 'SALTAR' (lo que ya existe no se
    toca) o 'ACTUALIZAR' (se pone como dice el Excel). Con errores no guarda nada."""
    resultado = {'creados': [], 'actualizados': [], 'saltados': [], 'sin_cambios': 0,
                 'repetidos': [], 'errores': [], 'simulado': simular, 'guardado': False}
    try:
        unidades, departamentos, resultado['errores'] = _leer(archivo)
    except (ValueError, KeyError, OSError) as error:
        resultado['errores'].append(str(error))
        return resultado
    except Exception:
        resultado['errores'].append('No se pudo leer el fichero: ¿es un Excel (.xlsx)?')
        return resultado
    if resultado['errores']:
        return resultado

    actualizar = estrategia == 'ACTUALIZAR'
    try:
        with transaction.atomic():
            por_nombre = {clave(u.descripcion): u for u in UnidadOrganizativa.objects.select_related('tipo')}
            # Primero las que no cuelgan de otra: así una subunidad ya encuentra a su principal.
            unidades.sort(key=lambda d: bool(d['tipo'].es_subunidad))
            for datos in unidades:
                padre = None
                if datos['tipo'].es_subunidad:
                    padre = por_nombre.get(clave(datos['padre']))
                    if padre is None:
                        resultado['errores'].append('Unidades, fila %d: la unidad principal «%s» no existe.'
                                                    % (datos['fila'], datos['padre']))
                        continue
                unidad = None
                if datos['codigo']:
                    unidad = UnidadOrganizativa.objects.filter(pk=datos['codigo']).first()
                if unidad is None and not datos['codigo']:
                    unidad = por_nombre.get(clave(datos['nombre']))
                nueva = unidad is None
                if nueva:
                    if not datos['tipo'].es_subunidad and not datos['grupo']:
                        resultado['errores'].append('Unidades, fila %d («%s»): una unidad nueva necesita su '
                                                    'grupo de nómina.' % (datos['fila'], datos['nombre']))
                        continue
                    unidad = UnidadOrganizativa(codigo_interno=datos['codigo'])
                cambios = []
                nuevos = {'descripcion': datos['nombre'], 'tipo': datos['tipo'], 'padre': padre,
                          'municipio': datos['municipio'] if datos['municipio'] else (None if nueva else unidad.municipio),
                          'grupo_nomina': datos['grupo'] or (None if nueva else unidad.grupo_nomina),
                          'orden_informe': datos['posicion'] or (None if nueva else unidad.orden_informe)}
                for campo, valor in nuevos.items():
                    if campo in ('tipo', 'padre', 'municipio'):
                        # Por id: una unidad nueva aún no tiene tipo y leerlo fallaría.
                        actual, deseado = getattr(unidad, campo + '_id'), getattr(valor, 'pk', None)
                    else:
                        actual, deseado = getattr(unidad, campo), valor
                    if actual != deseado:
                        cambios.append(ETIQUETAS[campo])
                if not nueva and not cambios:
                    resultado['sin_cambios'] += 1
                    continue
                if not nueva and not actualizar:
                    resultado['saltados'].append('Unidad %s: cambiaría %s' % (unidad.descripcion, ', '.join(cambios)))
                    continue
                anterior = unidad.descripcion
                for campo, valor in nuevos.items():
                    setattr(unidad, campo, valor)
                if unidad.orden_informe is None:
                    unidad.orden_informe = _siguiente(UnidadOrganizativa.objects.filter(padre=padre))
                otra = unidad_con_el_grupo_de_nomina(unidad.grupo_nomina, unidad.tipo, excluir_pk=unidad.pk)
                if otra is not None:
                    resultado['errores'].append(
                        'Unidades, fila %d («%s»): el grupo de nómina %s ya lo tiene «%s». Solo lo comparten '
                        'una unidad principal y sus subunidades.' % (datos['fila'], datos['nombre'],
                                                                     unidad.grupo_nomina, otra))
                    continue
                try:
                    unidad.save()
                except ValidationError as error:
                    resultado['errores'].append('Unidades, fila %d («%s»): %s' % (
                        datos['fila'], datos['nombre'], ' '.join(error.messages)))
                    continue
                por_nombre[clave(unidad.descripcion)] = unidad
                if nueva:
                    resultado['creados'].append('Unidad %s' % unidad.descripcion)
                else:
                    resultado['actualizados'].append('Unidad %s%s: %s' % (
                        anterior, ' (ahora «%s»)' % unidad.descripcion if anterior != unidad.descripcion else '',
                        ', '.join(cambios)))

            for datos in departamentos:
                unidad = por_nombre.get(clave(datos['unidad']))
                if unidad is None:
                    resultado['errores'].append('Departamentos, fila %d: la unidad «%s» no existe.'
                                                % (datos['fila'], datos['unidad']))
                    continue
                existente = next((d for d in Departamento.objects.filter(unidad_organizativa=unidad)
                                  if clave(d.descripcion) == clave(datos['nombre'])), None)
                if existente is None and unidad.tipo_id and unidad.tipo.es_principal:
                    resultado['errores'].append('Departamentos, fila %d («%s»): %s' % (
                        datos['fila'], datos['nombre'], MENSAJE_UNIDAD_PRINCIPAL % unidad.descripcion))
                    continue
                if existente is None:
                    Departamento.objects.create(
                        unidad_organizativa=unidad, descripcion=datos['nombre'],
                        orden_informe=datos['posicion'] or _siguiente(Departamento.objects.filter(
                            unidad_organizativa=unidad)))
                    resultado['creados'].append('Departamento %s (%s)' % (datos['nombre'], unidad.descripcion))
                    continue
                cambia = [ETIQUETAS[c] for c, v in (('descripcion', datos['nombre']),
                                                    ('orden_informe', datos['posicion']))
                          if v is not None and getattr(existente, c) != v]
                if not cambia:
                    resultado['sin_cambios'] += 1
                elif not actualizar:
                    resultado['saltados'].append('Departamento %s (%s): cambiaría %s' % (
                        existente.descripcion, unidad.descripcion, ', '.join(cambia)))
                else:
                    existente.descripcion = datos['nombre']
                    if datos['posicion'] is not None:
                        existente.orden_informe = datos['posicion']
                    existente.save()
                    resultado['actualizados'].append('Departamento %s (%s): %s' % (
                        existente.descripcion, unidad.descripcion, ', '.join(cambia)))

            if resultado['errores']:
                raise Simulacion()   # todo o nada
            # Las posiciones se comprueban al final (restricción diferida): forzarlo aquí
            # para poder explicar el error en vez de romper al confirmar.
            from django.db import connection
            with connection.cursor() as cursor:
                cursor.execute('SET CONSTRAINTS ALL IMMEDIATE')
            if simular:
                raise Simulacion()
        resultado['guardado'] = True
    except Simulacion:
        pass
    except IntegrityError:
        resultado['errores'].append(
            'Dos unidades de la misma rama (o dos departamentos de la misma unidad) quedarían con la '
            'misma Posición. Revise la columna «Posición» o déjela vacía para ponerlas al final. '
            'No se ha guardado nada.')
    if resultado['errores']:
        resultado['guardado'] = False
    return resultado


def exportar_estructura():
    """Las dos hojas, en el formato que lee `importar_estructura`."""
    unidades = [(u.codigo_interno, u.descripcion, u.tipo.descripcion,
                 u.padre.descripcion if u.padre_id else '', u.municipio.nombre if u.municipio_id else '',
                 u.grupo_nomina, u.orden_informe)
                for u in UnidadOrganizativa.objects.select_related('tipo', 'padre', 'municipio')
                .order_by('padre_id', 'orden_informe', 'descripcion')]
    departamentos = [(d.unidad_organizativa.descripcion, d.descripcion, d.orden_informe)
                     for d in Departamento.objects.select_related('unidad_organizativa')
                     .order_by('unidad_organizativa__orden_informe', 'unidad_organizativa__descripcion',
                               'orden_informe', 'descripcion')]
    libro = libro_vacio()
    hoja_con_titulos(libro, HOJA_UNIDADES, ['Código', 'Unidad', 'Tipo', 'Pertenece a', 'Municipio',
                                            'Grupo de nómina', 'Posición'],
                     unidades, [10, 48, 22, 30, 26, 16, 11])
    hoja_con_titulos(libro, HOJA_DEPARTAMENTOS, ['Unidad', 'Departamento', 'Posición'],
                     departamentos, [48, 60, 11])
    return libro
