# -*- coding: utf-8 -*-
"""El Word y el Excel de Informe Inteligente, con la imagen de Informes RRHH.

Los dos se construyen POR CÓDIGO, sin plantilla intermedia, y por el mismo
motivo: el número de columnas lo elige el usuario en el panel de campos. Una
plantilla `.docx` de docxtpl repite filas, no columnas, y una plantilla `.xlsx`
tendría que declarar de antemano cuántas columnas trae la tabla.

Los dos documentos llevan lo mismo y en el mismo orden, que es también el de la
pantalla:

    1. Quién informa (cabecera con logo) y de qué informe se trata.
    2. Los filtros aplicados. Sin ellos el papel es irrepetible: dos informes
       del mismo mes con condiciones distintas se ven exactamente igual.
    3. Las tarjetas de resumen que el usuario eligió.
    4. Los desgloses pedidos.
    5. La tabla de datos.
"""
import openpyxl
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import exportacion

# --------------------------------------------------------------------------
# Paleta, la misma de `static/css/informes-rrhh.css`
# --------------------------------------------------------------------------

AZUL = RGBColor(0x1E, 0x3A, 0x8A)
PIZARRA = RGBColor(0x1E, 0x29, 0x3B)
GRIS = RGBColor(0x47, 0x55, 0x69)
NARANJA = RGBColor(0xC2, 0x41, 0x0C)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)

HEX_AZUL = '1E3A8A'
HEX_PIZARRA = '1E293B'
HEX_GRIS = '475569'
HEX_SUAVE = 'F8FAFC'
HEX_BORDE = 'CBD5E1'


# ==========================================================================
# WORD
# ==========================================================================

def _escribir(parrafo, texto, tamano=11, negrita=False, color=PIZARRA):
    run = parrafo.add_run(str(texto))
    run.bold = negrita
    run.font.size = Pt(tamano)
    run.font.color.rgb = color
    return run


def _sombrear(celda, hexadecimal):
    """Color de fondo de una celda. python-docx no lo expone: hay que bajar al XML."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    sombra = OxmlElement('w:shd')
    sombra.set(qn('w:val'), 'clear')
    sombra.set(qn('w:fill'), hexadecimal)
    celda._tc.get_or_add_tcPr().append(sombra)


def documento_word_del_informe(context):
    """El informe en Word, apaisado y con la cabecera de siempre."""
    documento = Document()

    seccion = documento.sections[0]
    seccion.orientation = WD_ORIENT.LANDSCAPE
    seccion.page_width, seccion.page_height = seccion.page_height, seccion.page_width
    seccion.left_margin = seccion.right_margin = Cm(1.4)
    seccion.top_margin = seccion.bottom_margin = Cm(1.2)

    _cabecera_word(documento, seccion, context)
    _titulo_word(documento, context)
    _filtros_word(documento, context)
    _tarjetas_word(documento, context)
    _desgloses_word(documento, context)
    _tabla_word(documento, seccion, context)

    return documento


def _cabecera_word(documento, seccion, context):
    """Logo, empresa, territorio y quién exportó, repetidos en cada página.

    Van en el encabezado de Word y no en el cuerpo: así los repite el propio
    procesador de textos y no hay que duplicar nada al pasar de hoja.
    """
    cabecera = seccion.header
    cabecera.paragraphs[0].text = ''
    tabla = cabecera.add_table(rows=1, cols=2, width=seccion.page_width - Cm(2.8))
    izquierda, derecha = tabla.rows[0].cells

    config = context.get('config')
    if config and getattr(config, 'logo', None):
        try:
            izquierda.paragraphs[0].add_run().add_picture(config.logo.path,
                                                          width=Cm(2.4))
        except Exception:
            # Un logo que falte o esté corrupto no puede tumbar la exportación.
            pass

    parrafo = izquierda.add_paragraph()
    _escribir(parrafo, context['nombre_empresa'], 14, negrita=True, color=AZUL)

    territorio = ' · '.join(p for p in (context['cab_provincia'],
                                        context['cab_unidades_resumen']) if p)
    _escribir(izquierda.add_paragraph(), territorio, 9, negrita=True, color=PIZARRA)

    if context['cab_exportado_por']:
        _escribir(izquierda.add_paragraph(),
                  'Gestión de Capital Humano · Exportado por %s'
                  % context['cab_exportado_por'], 9, color=GRIS)

    parrafo = derecha.paragraphs[0]
    parrafo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _escribir(parrafo, context['periodo_texto'], 10, negrita=True, color=GRIS)

    pie = seccion.footer.paragraphs[0]
    _escribir(pie, 'Generado: %s' % context['fecha_generacion'], 9, color=GRIS)
    pie.add_run('\t\t')
    _escribir(pie, 'Orbith · Informe Inteligente', 9, color=GRIS)


def _titulo_word(documento, context):
    parrafo = documento.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _escribir(parrafo, context['titulo_informe'].upper(), 18, negrita=True,
              color=AZUL)

    # Nombre y descripción solo si el informe se cargó de los guardados: uno
    # armado sobre la marcha no tiene ninguno de los dos.
    if context['informe_nombre']:
        parrafo = documento.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _escribir(parrafo, context['informe_nombre'], 12, negrita=True,
                  color=PIZARRA)
    if context['informe_descripcion']:
        parrafo = documento.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _escribir(parrafo, context['informe_descripcion'], 10, color=GRIS)


def _filtros_word(documento, context):
    filtros = context['filtros_aplicados']
    parrafo = documento.add_paragraph()
    if filtros:
        _escribir(parrafo, 'Filtros aplicados: ', 10, negrita=True, color=AZUL)
        _escribir(parrafo, ' · '.join(filtros), 10, color=PIZARRA)
    else:
        _escribir(parrafo, 'Sin filtros: todos los registros.', 10, color=GRIS)


def _tarjetas_word(documento, context):
    indicadores = context['indicadores']
    if not indicadores:
        return

    tabla = documento.add_table(rows=2, cols=len(indicadores))
    for columna, indicador in enumerate(indicadores):
        arriba = tabla.cell(0, columna).paragraphs[0]
        arriba.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _escribir(arriba, indicador['label'].upper(), 8, negrita=True, color=GRIS)

        abajo = tabla.cell(1, columna).paragraphs[0]
        abajo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _escribir(abajo, indicador['valor'], 18, negrita=True, color=AZUL)

        _sombrear(tabla.cell(0, columna), 'F1F5F9')
        _sombrear(tabla.cell(1, columna), HEX_SUAVE)
    documento.add_paragraph()


def _desgloses_word(documento, context):
    for grupo in context['desgloses']:
        _escribir(documento.add_paragraph(),
                  'Desglose por %s' % grupo['etiqueta'].lower(), 12,
                  negrita=True, color=AZUL)

        tabla = documento.add_table(rows=1 + len(grupo['filas']), cols=3)
        tabla.style = 'Table Grid'
        for columna, titulo in enumerate(('Grupo', 'Cantidad', '%')):
            celda = tabla.cell(0, columna)
            parrafo = celda.paragraphs[0]
            parrafo.alignment = (WD_ALIGN_PARAGRAPH.LEFT if columna == 0
                                 else WD_ALIGN_PARAGRAPH.CENTER)
            _escribir(parrafo, titulo, 9, negrita=True, color=BLANCO)
            _sombrear(celda, HEX_AZUL)

        for numero, fila in enumerate(grupo['filas'], start=1):
            _escribir(tabla.cell(numero, 0).paragraphs[0],
                      fila['valor'] if fila['valor'] not in (None, '') else 'Sin definir',
                      10, color=PIZARRA)
            for columna, valor, color in ((1, fila['cantidad'], PIZARRA),
                                          (2, '%s %%' % fila['porcentaje'], NARANJA)):
                parrafo = tabla.cell(numero, columna).paragraphs[0]
                parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _escribir(parrafo, valor, 10, negrita=True, color=color)
        documento.add_paragraph()


#: Los dos campos cuyo valor lleva su color global en los documentos.
_CAMPOS_CON_COLOR = ('sexo', 'categoria_ocupacional')


def _color_de_celda(key, valor):
    """El color de la letra de una celda: el global del sexo o de la categoría.

    Es el mismo criterio que en la pantalla y en el Excel. La categoría llega
    ya como inicial —O, T, A, S, C—, así que se vuelve al código para poder
    consultar la paleta.
    """
    from nomencladores.models import color_de_sexo

    texto = str(valor or '').strip().upper()
    if key == 'sexo' and texto in ('F', 'M'):
        return _hex_a_rgb(color_de_sexo(texto))
    if key == 'categoria_ocupacional':
        codigo = {'O': 'OPE', 'T': 'TEC', 'A': 'ADM',
                  'S': 'SER', 'C': 'CDI'}.get(texto[:1])
        if codigo:
            from nomencladores.models import color_de_categoria
            return _hex_a_rgb(color_de_categoria(codigo))
    return PIZARRA


def _hex_a_rgb(hexadecimal):
    """«#db2777» → RGBColor. python-docx no acepta la cadena directamente."""
    limpio = (hexadecimal or '').lstrip('#')
    return RGBColor(int(limpio[0:2], 16), int(limpio[2:4], 16),
                    int(limpio[4:6], 16))


def _tabla_word(documento, seccion, context):
    columnas = context['campos_elegidos']
    if context['modo'] == 'resumen' or not columnas:
        _escribir(documento.add_paragraph(),
                  'Resumen sin listado: %s registros. Use la exportación a Excel '
                  'para el detalle completo.' % context['total_registros'],
                  10, negrita=True, color=NARANJA)
        return

    _escribir(documento.add_paragraph(), 'Detalle', 12, negrita=True, color=AZUL)

    tabla = documento.add_table(rows=1 + len(context['filas']), cols=len(columnas))
    tabla.style = 'Table Grid'
    tabla.autofit = False

    # El ancho de cada columna sale del reparto proporcional que calculó
    # `exportacion.repartir_anchos` sobre el contenido real.
    util = seccion.page_width - seccion.left_margin - seccion.right_margin
    for indice, columna in enumerate(columnas):
        celda = tabla.cell(0, indice)
        parrafo = celda.paragraphs[0]
        parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _escribir(parrafo, columna['label'], 8, negrita=True, color=BLANCO)
        _sombrear(celda, HEX_AZUL)
        for fila in tabla.rows:
            fila.cells[indice].width = int(util * columna['ancho'] / 100)

    for numero, valores in enumerate(context['matriz'], start=1):
        for indice, valor in enumerate(valores):
            columna = columnas[indice]
            parrafo = tabla.cell(numero, indice).paragraphs[0]
            # Los campos de contenido corto van centrados; los textos largos, a
            # la izquierda. Lo decide `metadata.es_campo_corto`, no esta función.
            if columna.get('corto'):
                parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _escribir(parrafo, '' if valor in (None, '') else valor, 8,
                      negrita=columna['key'] in _CAMPOS_CON_COLOR,
                      color=_color_de_celda(columna['key'], valor))

    if context['filas_omitidas']:
        _escribir(documento.add_paragraph(),
                  'Se muestran %s de %s filas. Use la exportación a Excel para el '
                  'listado completo.' % (context['filas_mostradas'],
                                         context['total_registros']),
                  9, negrita=True, color=NARANJA)


# ==========================================================================
# EXCEL
# ==========================================================================

LINEA = Side(style='thin', color=HEX_BORDE)
REJILLA = Border(left=LINEA, right=LINEA, top=LINEA, bottom=LINEA)

FUENTE_TITULO = Font(name='Calibri', size=14, bold=True, color=HEX_AZUL)
FUENTE_PORTADA = Font(name='Calibri', size=22, bold=True, color=HEX_AZUL)
FUENTE_CABECERA = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
FUENTE_DATO = Font(name='Calibri', size=10, color=HEX_PIZARRA)
FUENTE_ROTULO = Font(name='Calibri', size=9, bold=True, color=HEX_GRIS)
FUENTE_VALOR = Font(name='Calibri', size=12, bold=True, color=HEX_PIZARRA)
FUENTE_PIE = Font(name='Calibri', size=8, italic=True, color=HEX_GRIS)

RELLENO_CABECERA = PatternFill('solid', fgColor=HEX_AZUL)
RELLENO_ALTERNO = PatternFill('solid', fgColor=HEX_SUAVE)

IZQUIERDA = Alignment(horizontal='left', vertical='center', wrap_text=True)
CENTRADO = Alignment(horizontal='center', vertical='center', wrap_text=True)


def libro_excel_del_informe(context):
    """Portada, Resumen y Datos, como los libros de Informes RRHH."""
    libro = openpyxl.Workbook()
    libro.remove(libro.active)

    _portada_excel(libro, context)
    _resumen_excel(libro, context)
    _datos_excel(libro, context)
    return libro


def _pie_excel(hoja, fila, columnas, context):
    hoja.merge_cells(start_row=fila, start_column=1,
                     end_row=fila, end_column=max(columnas, 1))
    celda = hoja.cell(row=fila, column=1)
    celda.value = ' · '.join(p for p in (
        context['nombre_empresa'], context['cab_provincia'],
        context['cab_unidades_resumen'],
        'Exportado por %s' % context['cab_exportado_por'],
        context['fecha_generacion']) if p)
    celda.font = FUENTE_PIE
    celda.alignment = IZQUIERDA


def _portada_excel(libro, context):
    hoja = libro.create_sheet('Portada')
    hoja.sheet_view.showGridLines = False
    hoja.row_dimensions[1].height = 88

    config = context.get('config')
    if config and getattr(config, 'logo', None):
        try:
            from openpyxl.drawing.image import Image
            imagen = Image(config.logo.path)
            proporcion = (imagen.width / imagen.height) if imagen.height else 1
            imagen.height = 110
            imagen.width = int(110 * proporcion)
            hoja.add_image(imagen, 'A1')
        except Exception:
            pass

    hoja['B2'] = context['nombre_empresa']
    hoja['B2'].font = FUENTE_PORTADA
    hoja['B3'] = context['titulo_informe']
    hoja['B3'].font = FUENTE_TITULO

    datos = [('Provincia', context['cab_provincia'] or '—'),
             ('Alcance', context['cab_unidades_resumen']),
             ('Período', context['periodo_texto']),
             ('Exportado por', context['cab_exportado_por'] or '—'),
             ('Generado', context['fecha_generacion']),
             ('Registros', context['total_registros'])]
    if context['informe_nombre']:
        datos.insert(0, ('Informe guardado', context['informe_nombre']))
    if context['informe_descripcion']:
        datos.insert(1, ('Descripción', context['informe_descripcion']))

    fila = 6
    for bloque, titulo in ((datos, 'Identificación'),
                           ([(f, '') for f in context['filtros_aplicados']],
                            'Filtros aplicados')):
        if not bloque:
            continue
        hoja.cell(row=fila, column=2).value = titulo.upper()
        hoja.cell(row=fila, column=2).font = FUENTE_TITULO
        fila += 1
        for rotulo, valor in bloque:
            hoja.cell(row=fila, column=2).value = rotulo
            hoja.cell(row=fila, column=2).font = FUENTE_ROTULO
            hoja.cell(row=fila, column=3).value = valor
            hoja.cell(row=fila, column=3).font = FUENTE_VALOR
            hoja.cell(row=fila, column=3).alignment = IZQUIERDA
            fila += 1
        fila += 1

    for columna, ancho in ((1, 12), (2, 34), (3, 70)):
        hoja.column_dimensions[get_column_letter(columna)].width = ancho
    return hoja


def _resumen_excel(libro, context):
    """Las tarjetas elegidas y, debajo, los desgloses pedidos."""
    hoja = libro.create_sheet('Resumen')
    hoja.sheet_view.showGridLines = False

    hoja['B2'] = 'Resumen del informe'
    hoja['B2'].font = FUENTE_TITULO

    fila = 4
    for indicador in context['indicadores']:
        hoja.cell(row=fila, column=2).value = indicador['label']
        hoja.cell(row=fila, column=2).font = FUENTE_ROTULO
        celda = hoja.cell(row=fila, column=3)
        celda.value = indicador['valor']
        celda.font = Font(name='Calibri', size=16, bold=True, color=HEX_AZUL)
        celda.alignment = IZQUIERDA
        if indicador.get('porcentaje') is not None:
            hoja.cell(row=fila, column=4).value = '%s %%' % indicador['porcentaje']
            hoja.cell(row=fila, column=4).font = FUENTE_PIE
        fila += 1

    if not context['indicadores']:
        hoja.cell(row=fila, column=2).value = (
            'No se eligió ninguna tarjeta de resumen.')
        hoja.cell(row=fila, column=2).font = FUENTE_PIE
        fila += 1

    for grupo in context['desgloses']:
        fila += 1
        hoja.cell(row=fila, column=2).value = (
            'Desglose por %s' % grupo['etiqueta'].lower()).upper()
        hoja.cell(row=fila, column=2).font = FUENTE_TITULO
        fila += 1
        for columna, titulo in ((2, 'Grupo'), (3, 'Cantidad'), (4, '%')):
            celda = hoja.cell(row=fila, column=columna)
            celda.value = titulo
            celda.font = FUENTE_CABECERA
            celda.fill = RELLENO_CABECERA
            celda.alignment = CENTRADO if columna > 2 else IZQUIERDA
            celda.border = REJILLA
        fila += 1
        for detalle in grupo['filas']:
            valores = (detalle['valor'] if detalle['valor'] not in (None, '')
                       else 'Sin definir',
                       detalle['cantidad'], detalle['porcentaje'])
            for columna, valor in zip((2, 3, 4), valores):
                celda = hoja.cell(row=fila, column=columna)
                celda.value = valor
                celda.font = FUENTE_DATO
                celda.border = REJILLA
                celda.alignment = CENTRADO if columna > 2 else IZQUIERDA
            fila += 1

    for columna, ancho in ((1, 4), (2, 44), (3, 16), (4, 12)):
        hoja.column_dimensions[get_column_letter(columna)].width = ancho
    _pie_excel(hoja, fila + 1, 4, context)
    return hoja


def _datos_excel(libro, context):
    """La tabla, completa: el Excel no recorta nunca.

    Aunque en pantalla se haya pedido «solo el resumen» para el PDF, aquí va
    todo: es la salida a la que se invita cuando el listado no cabe en papel, y
    devolverla vacía dejaría al usuario sin ninguna forma de obtenerlo.
    """
    hoja = libro.create_sheet('Datos')
    columnas = context['campos_elegidos']

    hoja.cell(row=2, column=1).value = context['titulo_informe']
    hoja.cell(row=2, column=1).font = FUENTE_TITULO

    for indice, columna in enumerate(columnas, start=1):
        celda = hoja.cell(row=4, column=indice)
        celda.value = columna['label']
        celda.font = FUENTE_CABECERA
        celda.fill = RELLENO_CABECERA
        celda.alignment = CENTRADO
        celda.border = REJILLA
        # El ancho sale del mismo reparto que usa el PDF, escalado a caracteres.
        hoja.column_dimensions[get_column_letter(indice)].width = max(
            10, round(columna['ancho'] * exportacion.ANCHO_UTIL_EN_CARACTERES / 100))
    hoja.row_dimensions[4].height = 30

    for numero, valores in enumerate(context['matriz']):
        for indice, valor in enumerate(valores, start=1):
            columna = columnas[indice - 1]
            celda = hoja.cell(row=5 + numero, column=indice)
            celda.value = '' if valor is None else valor
            celda.border = REJILLA
            celda.alignment = CENTRADO if columna.get('corto') else IZQUIERDA

            # El sexo y la categoría llevan su color GLOBAL en la letra, igual
            # que en la pantalla y que en los Excel de RRHH. En la letra y no en
            # el fondo: una celda maciza en medio parte la lectura de la fila.
            if columna['key'] in _CAMPOS_CON_COLOR:
                color = _color_de_celda(columna['key'], valor)
                celda.font = Font(name='Calibri', size=10, bold=True,
                                  color='%02X%02X%02X' % (color[0], color[1], color[2])
                                  if isinstance(color, tuple) else str(color))
            else:
                celda.font = FUENTE_DATO

            if numero % 2:
                celda.fill = RELLENO_ALTERNO

    if columnas and context['matriz']:
        hoja.auto_filter.ref = 'A4:%s%d' % (
            get_column_letter(len(columnas)), 4 + len(context['matriz']))
    hoja.freeze_panes = hoja.cell(row=5, column=1)

    hoja.page_setup.orientation = 'landscape'
    hoja.page_setup.fitToWidth = 1
    hoja.sheet_properties.pageSetUpPr.fitToPage = True

    _pie_excel(hoja, 6 + len(context['matriz']), max(len(columnas), 1), context)
    return hoja
